from pathlib import Path
from collections import Counter
from zipfile import ZipFile
from lxml import etree as E
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
import json,re,shutil
T=Path(__file__).resolve().parent;R=T.parents[1];D=Path('/mnt/d/Paper/Thesis/Gemini')
source=D/'重构论文_格式恢复_20261016.docx'
refs=json.loads((T/'references.json').read_text());byref={r['number']:r for r in refs}
rows=[]
for filename in ['audit_001_073.json','audit_074_146.json','audit_147_218.json']:rows+=json.loads((T/filename).read_text())
rows.sort(key=lambda x:x['number']);assert [r['number']for r in rows]==list(range(1,219))
# Detect source changes during the audit: compare the actual reference text, not a hash.
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
with ZipFile(source)as z:root=E.fromstring(z.read('word/document.xml'))
live=[''.join(n.text or''for n in p.iter(W+'t'))for p in root.find(W+'body')if p.find(W+'pPr/'+W+'pStyle')is not None and p.find(W+'pPr/'+W+'pStyle').get(W+'val')=='bibliography']
assert live==[r['text']for r in refs],'Source references changed while checking'
status=Counter(r['status']for r in rows);unresolved=[]
for r in rows:
 n=r['number'];p=T/'doi_handles'/f'{n:03}.json';h=json.loads(p.read_text())if p.exists()else None
 if isinstance(r.get('canonical'),str):
  r['canonical_text']=r['canonical'];r['canonical']={k:r.get(k)for k in ['title','authors','year','venue','volume','issue','pages','identifier']if r.get(k)is not None}
 r['input_text']=byref[n]['text'];r['checked_date']='2026-10-18'
 if h and h.get('responseCode')==100:
  assert h.get('resolver',{}).get('http_status')==404
  r['doi_resolution']={'status':'not_found','doi':h['doi'],'http_status':404,'handle_responseCode':100,'sources':[h['source'],h['resolver']['url']]};unresolved.append(n)
 elif h and h.get('responseCode')==1:r['doi_resolution']={'status':'registered','doi':h['doi'],'source':h['source']}
 elif (T/'crossref'/f'{n:03}.json').exists():
  x=json.loads((T/'crossref'/f'{n:03}.json').read_text());assert x.get('http_status')==200;r['doi_resolution']={'status':'crossref_record','doi':x['queried_doi'],'source':x['source']}
 else:r['doi_resolution']={'status':'not_applicable'}
 assert r.get('sources'),n
summary={'date':'2026-10-18','source_document':str(source),'total':218,'bibliographic_status':dict(status),'doi_not_found_numbers':unresolved,'partially_unverified_numbers':[r['number']for r in rows if r['status']=='unverified'],'confirmed_bibliographic_mismatch_numbers':[r['number']for r in rows if r['status']=='correction_needed'],'source_references_unchanged':True,'word_modified':False}
(T/'audit_all_218.json').write_text(json.dumps({'summary':summary,'references':rows},ensure_ascii=False,indent=2))
def val(x):return '; '.join(str(v)for v in x)if isinstance(x,list)else str(x or'')
def issue_text(r):
 n=r['number']
 if n==122:return '正式出版题名、年份、卷期页码尚待直接出版记录核实；所列DOI在doi.org与Handle均未找到。'
 if n==139:return '作者次序、年份和卷期页码尚待直接原刊记录核实；DOI注册及原刊目标已确认。'
 if n==146:return '卷期页码尚待直接原刊记录核实；作者单位公开记录及DOI注册已确认。'
 if r['doi_resolution']['status']=='not_found':return '书目信息及所列DOI与原刊一致；doi.org与Handle当前未找到该标识。'
 return '\n'.join(i.get('finding',str(i))if isinstance(i,dict)else str(i)for i in r.get('issues',[]))

def urls(r):
 result=[s.get('url',str(s))if isinstance(s,dict)else str(s)for s in r['sources']]
 result+=r['doi_resolution'].get('sources',[])
 return list(dict.fromkeys(result))
label={'verified':'书目字段已核实','unverified':'部分字段未独立核实','correction_needed':'已确认需更正'}
dlabel={'not_found':'doi.org未找到（404）','registered':'DOI已注册','crossref_record':'DOI对应记录已核实','not_applicable':'无DOI／按URL或文献原件核对'}
wb=Workbook();ws=wb.active;ws.title='结论与待确认项'
for pair in [['核对日期','2026-10-18'],['核对Word',str(source)],['参考文献总数',218],['书目字段已核实',status['verified']],['部分字段未独立核实',status['unverified']],['已确认书目错配',status['correction_needed']],['当前DOI未找到',','.join(map(str,unresolved))],['论文文件','未修改'],['核对字段','题名、所列作者及次序、年份、出版物、卷期页码、版本和DOI/URL']]:ws.append(pair)
ws.append([]);ws.append(['待确认类型','编号','文献题名','具体记录','来源'])
for r in rows:
 if r['status']!='unverified'and r['doi_resolution']['status']!='not_found':continue
 typ='部分书目字段未核实'if r['status']=='unverified'else'DOI解析异常'
 if r['status']=='unverified'and r['doi_resolution']['status']=='not_found':typ+='；DOI解析异常'
 ws.append([typ,r['number'],r['canonical'].get('title',''),issue_text(r),'\n'.join(urls(r))])
allws=wb.create_sheet('218条逐条核验');headers=['编号','书目核验状态','Word原文','来源题名','来源作者','来源年份','来源出版物','卷','期','页码或文章号','DOI／arXiv／版本','DOI注册或解析','未核字段／异常说明','已核字段','证据链接','备注']
allws.append(headers)
for r in rows:
 c=r['canonical'];identifier=c.get('doi',c.get('identifier',c.get('arxiv','')))
 if not identifier:
  m=re.search(r'https://arxiv.org/abs/([^\s]+)',r['input_text']);identifier=m[1].rstrip('.')if m else''
 allws.append([r['number'],label[r['status']],r['input_text'],val(c.get('title')),val(c.get('authors')),val(c.get('year')),val(c.get('venue',c.get('container',c.get('publisher','')))),val(c.get('volume')),val(c.get('issue')),val(c.get('pages',c.get('page',c.get('article_number','')))),val(identifier),dlabel[r['doi_resolution']['status']],issue_text(r),val(r.get('verified_fields')), '\n'.join(urls(r)),val(r.get('notes'))])
 if urls(r)[0].startswith('http'):allws.cell(allws.max_row,15).hyperlink=urls(r)[0];allws.cell(allws.max_row,15).font=Font(name='等线',size=10,color='0563C1',underline='single')
for sheet,widths in [(ws,[26,18,54,88,80]),(allws,[7,24,100,65,42,12,55,9,9,19,36,33,90,65,100,90])]:
 for i,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(i)].width=width
 for row in sheet:
  for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
 for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='244062');cell.font=Font(name='等线',size=11,color='FFFFFF',bold=True)
 sheet.freeze_panes='A2'
allws.auto_filter.ref=allws.dimensions
for i,r in enumerate(rows,2):
 allws.row_dimensions[i].height=90
 if r['status']=='unverified'or r['doi_resolution']['status']=='not_found':
  for cell in allws[i]:cell.fill=PatternFill('solid',fgColor='FFF2CC')
name='参考文献逐条核验_20261018';out=R/(name+'.xlsx');wb.save(out)
# Re-open to catch omitted rows or altered reference numbering in the delivered table.
check=load_workbook(out,read_only=True,data_only=True);assert check['218条逐条核验'].max_row==219;assert [check['218条逐条核验'].cell(i,1).value for i in range(2,220)]==list(range(1,219));check.close()
lines=[f'核对对象：{source}', '核对日期：2026-10-18。核对题名、所列作者及次序、年份、出版物、卷期页码、版本及DOI/URL。', '', f'共218条：{status["verified"]}条书目字段与来源匹配；{status["unverified"]}条仍有部分字段未取得直接出版记录的独立复核。未确认存在虚构文献或实质书目错配。', '', '目前不能给出“218条全部准确无误”的无条件结论。', '', '当前DOI解析异常：'+ '、'.join('['+str(n)+']'for n in unresolved)+'。这些标识与原刊/已有书目所列对应，但doi.org与Handle均返回未找到；与文献真实性分开记录。', '', '| 编号 | 待确认事项 | 来源 |', '|---|---|---|']
for r in rows:
 if r['status']=='unverified'or r['doi_resolution']['status']=='not_found':lines.append('| ['+str(r['number'])+'] | '+issue_text(r).replace('\n','；').replace('|','\\|')+' | '+'；'.join('[来源'+str(i+1)+']('+u+')'for i,u in enumerate(urls(r))if u.startswith('http'))+' |')
lines+=['','139、146的DOI注册及原刊目标已确认；其余未核字段详见表中记录。论文、Markdown及参考文献原文均未修改。','',f'218条的逐条元数据与来源链接见同名Excel，机器可读记录见[完整核验记录]({T}/audit_all_218.json)。']
(R/(name+'.md')).write_text('\n'.join(lines)+'\n')
for suffix in ['.xlsx','.md']:shutil.copy2(R/(name+suffix),D/(name+suffix))
print(json.dumps(summary,ensure_ascii=False));print(out)
