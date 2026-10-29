from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
from lxml import etree as E
from docx import Document
import json

A = Path(__file__).resolve().parent
SRC = A.parents[1] / '调整完格式的.docx'
OLD = Path('/mnt/d/Paper/Thesis/Gemini/重构论文_导言小结同步_20261018.docx')
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'

def row(path, label):
    with ZipFile(path) as z:
        root = E.fromstring(z.read('word/document.xml'))
    return next(tr for tr in root.iter(W+'tr') if label in ''.join(t.text or '' for t in tr.iter() if t.tag in (W+'t', M+'t')))

cur = row(SRC, '(E-46)')
old = row(OLD, '(E-46)')
cases = []
cases.append(('A - Current equation, unchanged', deepcopy(cur)))
b = deepcopy(cur)
for r in b.iter(M+'r'):
    if any('\u200b' in (t.text or '') for t in r.iter(M+'t')):
        for n in list(r.findall(M+'rPr/'+M+'nor')):
            n.getparent().remove(n)
cases.append(('B - Normal-text property removed only from four placeholders', b))
c = deepcopy(cur)
for t in c.iter(M+'t'):
    if t.text:
        t.text = t.text.replace('\u200b', '')
cases.append(('C - Four zero-width characters removed, other properties unchanged', c))
d = deepcopy(cur)
newmath = list(d.iter(M+'oMath'))
oldmath = list(old.iter(M+'oMath'))
assert len(newmath) == len(oldmath)
for nm, om in zip(newmath, oldmath):
    nm.getparent().replace(nm, deepcopy(om))
cases.append(('D - Previous native-math formatting, same mathematical content', d))
e = deepcopy(d)
for t in e.iter(M+'t'):
    if t.text:
        t.text = t.text.replace('\u200b', '')
cases.append(('E - Previous math formatting plus removal of four placeholders', e))

doc = Document()
section = doc.sections[0]
with ZipFile(SRC) as z:
    srcparts = {n:z.read(n) for n in ('word/styles.xml', 'word/settings.xml')}
    root = E.fromstring(z.read('word/document.xml'))
    sect = root.find(W+'body/'+W+'sectPr')
    if sect is not None:
        for n in list(section._sectPr):
            section._sectPr.remove(n)
        for n in sect:
            if n.tag not in (W+'headerReference', W+'footerReference'):
                section._sectPr.append(deepcopy(n))
for i, (label, tr) in enumerate(cases):
    p = doc.add_paragraph(label)
    p.paragraph_format.space_after = 20
    tbl = E.Element(W+'tbl')
    origtable = cur.getparent()
    for tag in ('tblPr', 'tblGrid'):
        node = origtable.find(W+tag)
        if node is not None:
            tbl.append(deepcopy(node))
    tbl.append(tr)
    for node in list(tbl.iter()):
        if node.tag in (W+'bookmarkStart', W+'bookmarkEnd', W+'lastRenderedPageBreak'):
            node.getparent().remove(node)
    p._p.addnext(tbl)
    if i < len(cases)-1:
        doc.add_page_break()
sourcebody = root.find(W+'body')
newbody = doc._element.find(W+'body')
sourcebody.clear()
for child in newbody:
    sourcebody.append(deepcopy(child))
with ZipFile(SRC) as zin, ZipFile(A/'E46_display_probes.docx', 'w', ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        zout.writestr(item, E.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True) if item.filename == 'word/document.xml' else zin.read(item.filename))
(A/'probe_cases.json').write_text(json.dumps([x[0] for x in cases], indent=2)+'\n')
print(str(A/'E46_display_probes.docx'))
