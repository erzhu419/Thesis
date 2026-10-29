"""Create a standalone, copyable achievements list; existing theses stay read-only."""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

A=Path(__file__).resolve().parent
R=A.parents[1]
OUT=R/'攻读学位期间主要的研究成果_20261019.docx'
arxiv=json.loads((A/'arxiv_metadata.json').read_text())['items']
local=json.loads((A/'local_cfcmt_or.json').read_text())['papers']
freq=json.loads((A/'local_freq_papers.json').read_text())['papers']

published=[
 {'title':'Single agent robust deep reinforcement learning for bus fleet control',
  'authors':['ZHANG Y','ZHENG L','ZHANG Q','TANG H'],
  'journal':'Transportation Safety and Environment','year':2026,'volume':'8','issue':'2','article':'tdag005','doi':'10.1093/tse/tdag005',
  'source':'https://academic.oup.com/tse/article/8/2/tdag005/8530551','position':1,'status':'已发表'},
 {'title':'Offline and offline-to-online reinforcement learning for bus holding control',
  'authors':['ZHANG Y','LI H','ZHENG L','ZHANG Q'],
  'journal':'Transportation Safety and Environment','year':2026,'article':'tdag059','doi':'10.1093/tse/tdag059',
  'source':'https://academic.oup.com/tse/advance-article/doi/10.1093/tse/tdag059/8812242','position':1,'status':'在线发表'},
 {'title':'TransitDuet: bi-level reinforcement learning for joint bus target-headway planning and holding control',
  'authors':['ZHANG Y','LI H','ZHENG L','ZHANG Q'],
  'journal':'Transportation Safety and Environment','position':1,'status':'已录用',
  'source':'用户提供题名、已录用状态，并明确作者信息与tdag059一致'},
 {'title':'Safety-oriented decision-making for connected and automated vehicles in highway mixed-autonomy traffic via dynamic attention-enhanced graph reinforcement learning',
  'authors':['LI H','ZHANG Y','ZHENG L','ZHANG Z'],
  'journal':'Transportation Safety and Environment','year':2026,'article':'tdag062','doi':'10.1093/tse/tdag062',
  'source':'https://academic.oup.com/tse/advance-article/doi/10.1093/tse/tdag062/8907179','position':2,'status':'在线发表'},
]
submitted=[]
for item in arxiv:
 submitted.append({'title':item['title'],'authors':item['official_citation_authors'],'year':item['year'],'arxiv_id':item['arxiv_id'],'url':item['source_url'],'source':item['version_url'],'position':1,'status':'正在投稿'})
or_paper=next(p for p in local if 'Operations Research' in p['project'])
cfcmt=next(p for p in local if p['project']=='CFCMT')
submitted.append({'title':or_paper['title'],'journal':'Operations Research','source':or_paper['manuscript'],'position':1,'status':'正在投稿'})
submitted.append({'title':cfcmt['title'],'source':cfcmt['manuscript'],'position':1,'status':'正在投稿'})
for p in freq:
 submitted.append({'title':p['title'],'source':p['sources'][0]['path'],'position':1,'status':'正在投稿'})
assert len(published)==4 and len(submitted)==9
assert sum(p['position']==2 for p in published+submitted)==1

def citation(item,number):
 authors=', '.join(item.get('authors',[]))
 start=f'[{number}] '+(authors+'. ' if authors else '')+item['title']
 if item['status']=='已录用':
  return start+'[J]. '+item['journal']+'（已录用）。（本人第一作者）'
 if item['status'] in ('已发表','在线发表'):
  typ='[J]' if 'volume' in item else '[J/OL]'
  volume=f', {item["volume"]}({item["issue"]})' if 'volume' in item else ''
  role='第二' if item['position']==2 else '第一'
  return start+typ+'. '+item['journal']+', '+str(item['year'])+volume+': '+item['article']+'. DOI: '+item['doi']+'。（本人'+role+'作者）'
 if 'arxiv_id' in item:
  return start+'[EB/OL]. arXiv:'+item['arxiv_id']+', '+str(item['year'])+'. '+item['url']+'。（本人第一作者）'
 return start+'. '+('Operations Research（投稿中）。' if item.get('journal')=='Operations Research' else '')+'（本人第一作者）'

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.8);sec.bottom_margin=Cm(2.8)
sec.left_margin=Cm(3);sec.right_margin=Cm(2.5)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
normal.paragraph_format.line_spacing=Pt(20)
normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(6)
expected=[]
def para(value,kind='body'):
 p=doc.add_paragraph();r=p.add_run(value);expected.append(value)
 r.font.name='Times New Roman';r.font.size=Pt(12);r.bold=False
 r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
 if kind=='title':
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r.font.size=Pt(16);r.bold=True
  r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体')
  p.paragraph_format.space_after=Pt(18);p.paragraph_format.keep_with_next=True
 elif kind=='heading':
  r.font.size=Pt(14);r.bold=True
  r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体')
  p.paragraph_format.space_before=Pt(10);p.paragraph_format.space_after=Pt(8)
  p.paragraph_format.keep_with_next=True
 else:
  p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  p.paragraph_format.left_indent=Cm(.75);p.paragraph_format.first_line_indent=Cm(-.75)
  p.paragraph_format.keep_together=True
 return p

para('攻读学位期间主要的研究成果','title')
para('一、发表的学术论文','heading')
for i,item in enumerate(published,1):para(citation(item,i))
para('二、正在投稿的学术论文','heading')
for i,item in enumerate(submitted,1):para(citation(item,i))
doc.core_properties.title='攻读学位期间主要的研究成果'
doc.core_properties.subject='已发表、已录用及正在投稿的学术论文'
doc.core_properties.author='Yifan Zhang'
doc.save(OUT)

W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
with ZipFile(OUT)as z:root=E.fromstring(z.read('word/document.xml'))
actual=[''.join(t.text or '' for t in p.iter(W+'t'))for p in root.find(W+'body')if p.tag==W+'p']
assert actual==expected
joined='\n'.join(actual)
assert joined.count('本人第一作者')==12 and joined.count('本人第二作者')==1
assert all(joined.count(p['title'])==1 for p in published+submitted)
assert 'Anonymous' not in joined and 'Author names withheld' not in joined
manifest={'output':str(OUT),'publication_metadata_sources':'出版社官网；arXiv链接使用当前官方题名和作者；本地匿名稿只列已核实题名，不补猜作者或期刊','published_or_accepted':published,'under_submission':submitted,'status_and_position_source':'用户提供录用/投稿状态及作者位次；tdag062第二作者也与官网署名一致','unmodified_thesis_documents':True,'paragraphs':expected}
(A/'achievements_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
(R/'攻读学位期间主要的研究成果_20261019.md').write_text('\n\n'.join(expected)+'\n')
print(json.dumps({'word':str(OUT),'published_or_accepted':len(published),'under_submission':len(submitted),'first_author':12,'second_author':1,'paragraphs':len(expected)},ensure_ascii=False))
