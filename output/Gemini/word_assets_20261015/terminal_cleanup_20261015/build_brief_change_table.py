"""从终审定稿记录生成短修改表；只写报告，不修改论文。"""
import json
import re
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
HEAD = ['位置', '短操作', '短旧→新']


def read(name):
    return json.loads((BASE / name).read_text())


def compact(text, limit=82):
    text = re.sub(r'\s+', ' ', text).strip().replace('**', '')
    return text if len(text) <= limit else text[:limit - 1] + '…'


original_lines = (BASE / '修改前_重构论文.md').read_text().splitlines()
journal = read('source_edit_journal.json')


def location(edit):
    lines = edit.get('original_md_lines', [])
    if 'start_line' in edit:
        suffix = f"原MD L{edit['start_line']}–{edit['end_line']}"
    elif 'line' in edit:
        suffix = f"原MD L{edit['line']}"
    else:
        if not lines:
            lines = [n for n, line in enumerate(original_lines, 1) if edit['old'] in line]
        if edit.get('id') == 'controlled-summary-sequence-2':
            lines = [2513]
        if edit.get('id') == 'controlled-intro-3':
            lines = [178]
        if edit.get('id') == 'G-reference':
            lines = [n for n in lines if n != 1991]  # 此处已由A9-1单列。
        if edit.get('id') == 'B10-1':
            lines = [n for n in lines if n != 2080]  # 此原词已随受控从句删除。
        suffix = '原MD L' + '、'.join(map(str, dict.fromkeys(lines))) if lines else '原MD'
    label = edit.get('location') or edit.get('id') or edit.get('issue', '文字')
    if not edit.get('location') and lines:
        headings = []
        for target in dict.fromkeys(lines):
            heading = next((z.lstrip('#').strip() for z in reversed(original_lines[:target]) if re.match(r'^#{1,4}\s', z)), '')
            match = re.match(r'([1-7](?:\.\d+)*|附录[A-G]|摘\s*要)', heading)
            headings.append(match.group(1) if match else '摘要/前言')
        sections = list(dict.fromkeys(headings))
        scope = '、'.join(sections) if len(sections) <= 3 else '多章术语'
        label = scope + ' / ' + label
    count = edit.get('count_applied', 1)
    return f"{label}（{suffix}{f'；{count}处' if count > 1 else ''}）"


large = {
    1: ('删除整节', '6.4.1受控动力学基准对照实验、6.4.2模型选择与组成因素消融分析→删除；含图6-3—6-9、表6-2—6-4'),
    2: ('删除专用定义', '受控固定时间步目标及追加目标交互定义（含式6-2）→删除'),
    3: ('删除专用示例', '受控连续增量预测目标示例（含式6-10）→删除'),
    4: ('删除专用过滤', '受控有效性过滤与低置信样本剔除（含式6-13）→删除'),
    5: ('删除附录小段', '原附录F（一）概率集成、（二）SINDy残差、（三）在线核加权、（四）训练恢复→删除；保留公交设置并同步为附录G'),
    6: ('删除受控结论', '原6.4.6第一项受控三任务数值证据→删除'),
    7: ('删除受控从句', '受控动力学平台上的追加目标交互实验从句→删除；公交任务句保留'),
    13: ('删除受控划分', '受控任务80%/10%/10%样本划分→删除；公交按运行日划分保留'),
    15: ('删除受控说明', '受控SINDy残差、近期状态核加权说明及其专用引文→删除'),
    16: ('删除受控示例', '受控γΔ=0.5、不动点0.25示例→删除；长时域公交说明保留'),
    17: ('删除受控引线', '两类实验及受控四问题引线→删除；保留多线路公交跨保真实验引线'),
    20: ('删除受控结论', '执行器偏移回报+1,529.9及其置信区间结论→删除；公交方法小结保留'),
    21: ('删除专用变量', '受控预测量ẑ及逐维标准差定义→删除；候选类别m与验证误差Rm保留'),
    23: ('删除专用公式', '式(6-11)：受控连续量标准化均方根误差→删除'),
    54: ('段末加短句', '原段末→追加“自动选择组比无模型基线多耗时约1.98分钟，单状态推理仍为毫秒量级。”'),
    55: ('删未用自定义', '历史运行数据及其括号内自定义→目标域固定离线数据'),
}


def snippet(edit):
    old, new = edit['old'], edit['new']
    if old and new.startswith(old) and len(new) > len(old):
        return '原句末→追加“' + compact(new[len(old):], 110) + '”'
    return compact(old) + '→' + (compact(new) if new else '删除')


source_rows = []
for number, edit in enumerate(journal, 1):
    if number in large:
        action, detail = large[number]
    else:
        detail = snippet(edit)
        action = '删除短片段' if not edit['new'] else '原位微调'
        if edit.get('id', '').startswith('B12'):
            action = '统一四级标题编号'
        elif edit.get('id') == 'B10-1':
            action = '统一术语（18处）'
        elif edit.get('issue') == 'A8':
            action = '补置信区间方向'
        elif edit.get('issue') == 'A7':
            action = '补回合计数因子'
    source_rows.append([location(edit), action, detail])


def word_location(edit):
    if edit.get('location'):
        return edit['location']
    index = edit['index']
    if index == 42:
        return '主要符号说明（原Word块42）'
    if index == 750:
        return '表4-2题注（原Word块750）'
    sec = 'F.1' if index < 2147 else 'F.2' if index < 2155 else 'F.3' if index < 2164 else 'F.4'
    return f'附录{sec}（原Word块{index}）'


word_rows = []
word_edits = read('word_only_edits.json')
for edit in word_edits['edits']:
    action = '删完成率内容' if not edit['new'] else '原位同步'
    if edit.get('index') == 2170:
        action = '删保障并修正回合数'
    elif edit.get('index') == 750:
        action = '删重复千位单位'
    elif edit.get('index') == 2148:
        action = '补章节定位、删重复字'
    word_rows.append([word_location(edit), action, snippet(edit)])
for edit in word_edits.get('row_deletions', []):
    word_rows.append([edit['location'], '删除仅供受控实验使用的符号行', compact(edit['old'], 110) + '→删除整行'])


figmap = read('renumbering.json')['figures']
figure_rows = []
for fig in read('manifest_ch1_4.json')['figures']:
    if fig.get('status') != 'modified':
        continue
    changes, details = fig.get('changes', []), []
    if fig.get('labels'):
        details.append('子图标题上方→对应子图下方')
    if any(x['type'] == 'color' for x in changes):
        details.append('调整区分度不足的系列颜色')
    if any(x['type'] == 'line_style' for x in changes):
        details.append('K=40实线→虚线；图例同步')
    if any(x['type'] == 'marker' for x in changes):
        details.append('每对两个空心圆→空心圆与菱形')
    if any(x['type'] == 'legend_title_collision' for x in changes):
        details.append('分开顶部图例与(a)标题')
    if any(x['type'] == 'minus_font' for x in changes):
        details.append('补全减号字形')
    if any(x['type'] == 'overlap_readability' for x in changes):
        details.append('重合曲线增加线型/圆方三角标记区分')
    if fig['figure'] == '图3-13':
        details += ['启用减关闭/关闭减启用→带减号并注明正值改善', '行动者/评论家→策略/价值；快速切换→运行模式快速切换', '补齐(b)场景文字、增加留白']
    figure_rows.append([fig['figure'], '图件修正', '；'.join(details)])
for fig in read('figures/manifest_56F.json')['figures']:
    label = fig['figure']
    oldnumber = label.removeprefix('图')
    if oldnumber in figmap and oldnumber != figmap[oldnumber]:
        label += '（现图' + figmap[oldnumber] + '）'
    detail = '；'.join(fig.get('actions', []))
    detail = detail.replace('原子图字母和原中文题完整移至该横轴标题下方；保留曲线、误差条、数值、字体与配色。', '子图字母和标题上方→对应横轴标题下方')
    figure_rows.append([label, '图件修正', compact(detail, 170)])

number_rows = [
    ['第6章图号', '删除后顺延', '原图6-10—6-18→图6-3—6-11'],
    ['第6章表号', '删除后顺延', '原表6-5—6-10→表6-2—6-7'],
    ['第6章公式号', '删除后顺延', '原式6-3—6-9→6-2—6-8；原式6-12→6-9；原式6-14—6-19→6-10—6-15'],
    ['第6章小节号', '删除后顺延', '原6.4.3—6.4.6→6.4.1—6.4.4；正文交叉引用同步'],
    ['附录定位', '同步既有Word附录号', '原MD附录F第五部分→附录G；新增附录F.2短引线，保留原Word附录F'],
    ['参考文献序号与链接', '删除后顺序重排', '226条→218条；保留条目的内部引用按原身份对应'],
]
ref_rows = [['参考文献：' + key, '删除仅由删文使用的条目', key + '→删除（保留内容中无引用）'] for key in read('citation_usage_after_controlled_deletion.json')['only_deleted_keys']]
preserve_rows = [
    ['修改前Markdown', '原件保留', str(BASE / '修改前_重构论文.md')],
    ['修改前Word', '原件保留', str(BASE / '修改前_通篇校核版.docx')],
]
unit_sync = read('md_unit_sync.json')
sync_rows = [[unit_sync['location'], 'MD同步既有Word千位显示', unit_sync['old'] + '→' + unit_sync['new'] + '；数值不变']]
sections = [('文字微调与删除', source_rows), ('Word独有修改', word_rows), ('MD单位同步', sync_rows), ('图件修正', figure_rows), ('编号与定位', number_rows), ('参考文献删除', ref_rows), ('原件保留路径', preserve_rows)]
assert len(figure_rows) == 41
assert len(ref_rows) == 8
assert len({row[0] for row in figure_rows}) == 41

md = ['# 终审修改简表', '', f'依据已定修订记录：正文{len(source_rows)}项操作、Word独有{len(word_rows)}项操作、MD单位同步1项、图件41张（第1–4章22张，第5–6章及附录F 19张），删除仅由删文引用的参考文献8条。操作项数不等于改动段落数。', '', '保留段落按原位微调，未拆分；原文件保留。下表只列受改片段，大段删除仅列标题及范围。位置和图表公式旧号以修改前文件为准；现编号见编号表。', '']
for title, rows in sections:
    md += [f'## {title}（{len(rows)}行）', '', '| ' + ' | '.join(HEAD) + ' |', '| --- | --- | --- |']
    for row in rows:
        md.append('| ' + ' | '.join(str(x).replace('|', '\\|').replace('\n', ' ') for x in row) + ' |')
    md.append('')
stem = '终审修改简表_20261015'
(BASE / (stem + '.md')).write_text('\n'.join(md))

book = Workbook()
book.remove(book.active)
for title, rows in sections:
    sheet = book.create_sheet(title)
    sheet.append(HEAD)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for n, width in enumerate([54, 25, 106], 1):
        sheet.column_dimensions[get_column_letter(n)].width = width
    for cell in sheet[1]:
        cell.font = Font(name='微软雅黑', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24476B')
    sheet.row_dimensions[1].height = 25
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            cell.font = Font(name='微软雅黑', size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.row_dimensions[cells[0].row].height = min(80, max(31, 16 * max((len(str(c.value)) + w - 1) // w for c, w in zip(cells, [30, 14, 57]))))
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:1'
book.save(BASE / (stem + '.xlsx'))
print(json.dumps({'files': [str(BASE / (stem + ext)) for ext in ['.md', '.xlsx']], 'rows': {title: len(rows) for title, rows in sections}, 'total_rows': sum(len(rows) for _, rows in sections)}, ensure_ascii=False))
