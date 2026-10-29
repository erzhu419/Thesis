"""Compact review artifacts for the requested symbol-table-only change."""
from pathlib import Path
from collections import Counter
import json
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

A = Path(__file__).resolve().parent
OUT = A.parents[1]
plan = json.loads((A / 'symbol_table_plan.json').read_text())
actions = {'keep': '保留', 'update': '按实际改写', 'delete': '删除未用行', 'add': '补收实际符号'}

def readable(tex):
    if not tex:
        return ''
    greek = {'alpha':'α','beta':'β','gamma':'γ','Gamma':'Γ','delta':'δ','Delta':'Δ',
             'epsilon':'ϵ','varepsilon':'ε','theta':'θ','lambda':'λ','mu':'μ','nu':'ν',
             'xi':'ξ','pi':'π','rho':'ρ','sigma':'σ','tau':'τ','phi':'ϕ','varphi':'φ',
             'chi':'χ','psi':'ψ','Psi':'Ψ','kappa':'κ','zeta':'ζ','Omega':'Ω'}
    s = tex.replace('$', '').replace('\\\\', '\\')
    for macro, mark in [('widehat','̂'),('hat','̂'),('overline','̄'),('bar','̄'),('widetilde','̃'),('tilde','̃')]:
        s = re.sub(r'\\' + macro + r'\s*\{?\s*([A-Za-z])\}?', lambda m:m[1]+mark, s)
    for name, value in greek.items():
        s = re.sub(r'\\' + name + r'\b', value, s)
    s = re.sub(r'\\(?:mathcal|mathsf|mathbb|mathbf|boldsymbol|mathrm|operatorname)\s*', '', s)
    s = s.replace('\\prime', '′').replace('\\star', '⋆').replace('\\min', 'min')
    s = s.replace('\\det', 'det').replace('\\max','max').replace('\\mid','∣')
    s = s.replace('\\odot','⊙').replace('\\parallel','∥').replace('\\cdot','·')
    s = s.replace('\\|', '∥').replace('\\dim','dim').replace('\\,','')
    return s.replace('{','').replace('}','').strip()

wb = Workbook()
sheet = wb.active
sheet.title = '逐项修改对照'
sheet.append(['原行号','处理','原符号','修订后符号','原含义','修订后含义','依据或范围'])
for item in plan['rows']:
    sheet.append([item['source_row'] or '新增', actions[item['action']],
                  readable(item['original_symbol']), readable(item['correct_tex']),
                  item['original_definition'], item['correct_definition'] or '', item['reason']])

scope = wb.create_sheet('同形符号作用域')
scope.append(['符号','实际不同用法','本轮处理'])
for row in [
    ('ν','第2章CVaR尾部质量；第5章状态价值参数','表中分别说明，公式保留'),
    ('ξ','第2章CVaR成本阈值；第6章置信网络参数','补置信参数及目标参数横线，注明作用域'),
    ('ζ','第2章CVaR重加权密度；第4章线段ζ(u)','删除原高斯潜变量释义，按实际改写'),
    ('G','乘客广义时间成本；第4章特征网络G','补特征网络条目，注明与成本G的区别'),
    ('τ、τ_h','积分时刻、目标网络软更新系数；奖励衰减尺度τ_h在实验中简写τ','表内区分作用域，正文17处简写保持原样'),
    ('τ_e','第2章IQL期望分位水平；第6章模型预测误差尺度','分别注明对应章节'),
    ('B','训练小批量；附录E.5拟合映射的Lipschitz常数','去掉未用的batch别名，注明局部复用'),
    ('κ_k、κ′_k、κ_ale','第3章目标范数；第4章当前/目标范数；附录A加权固定量','按实际完整记号和章节区分'),
    ('m、m_grp、M','类别数、同时到站待决策车辆数、候选动作数','替换旧的cat/batch/act专名'),
    ('I、1','正文双线体I示性函数；附录粗体1示性函数','保留两种实际字体，不误判为同一个字形'),
]:
    scope.append(row)

notes = wb.create_sheet('表外记录')
notes.append(['位置','检查结果','本轮处理'])
notes.append(['式3-11、3.4.5节及表3-6/3-7',
              '同一奖励时间尺度在正式式中为τ_h，实验说明共17处写τ；与软更新τ需按上下文区分。',
              '已在符号表注明别名及作用域，正文与公式未改。'])
notes.append(['2.7.3节与式3-32',
              '物理动作a和归一化动作ã的定义已给出；3-32中回放历史动作a是否已转为网络坐标，原文未单独说明。此项不等于实现算错。',
              '仅作待澄清记录，不改原文、不推断实现。'])

for ws in wb:
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(name='宋体', bold=True, color='FFFFFF', size=11)
        cell.fill = PatternFill('solid', fgColor='24455D')
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='宋体', size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[row[0].row].height = 58
    widths = [10,16,31,36,44,62,58] if ws is sheet else [26,86,68]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
path = OUT / '符号表修改对照_20261019.xlsx'
wb.save(path)

counts = Counter(item['action'] for item in plan['rows'])
text = f'''# 符号表修订记录

以《调整完格式的_公式符号复核_20261019.docx》实际公式为准，逐项核对原98条符号说明，覆盖正文及附录A–F。本轮只修符号说明区域，原件与MD保留。

原98条中保留{counts['keep']}条、改写{counts['update']}条、删除{counts['delete']}条，另补收{counts['add']}条，现共103条。部分删除未用符号的条目计入“改写”。完整前后对照见《符号表修改对照_20261019.xlsx》。

| 项目 | 修订 |
| --- | --- |
| 人数与乘客时间 | 按实际改为 $N^w(t),N^b(t)$、$T_p^w,T_p^b$。 |
| 目标网络 | 采用正文的 $\\phi_k'$、$\\widehat Q_{{\\phi_k'}}$，修正表前“统一负号”的旧说明。 |
| BAPR与CS-BAPR | 补入实际诊断量、运行长度权重、动态保守系数、导数损失及 $\\Delta d$。 |
| MC-WM | 对齐模型输出、置信网络、预测误差、样本权重及递推目标的真实记号。 |
| 无对应旧名 | 删除机器精度、模型特征维数、符号检验、假设观测归一化等未实际使用的整行；其余行只删除未用部分。 |
| 同形复用 | 按章节说明 $\\nu,\\xi,\\zeta,G,\\tau,B,\\kappa_k$ 等的不同用法，保留已有局部定义。 |

表外正文、4634个数学对象、手工空行、等号分段与53个图片部件均保持不变。式3-32的动作坐标约定另作记录，未擅自补写正文或推断实现。
'''
(OUT / '符号表修订记录_20261019.md').write_text(text, encoding='utf-8')
print(path)
