from pathlib import Path
import json
import re
import zipfile

import fitz
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

HERE = Path(__file__).resolve().parent
OUT = HERE.parents[1] / 'output'
PATH = OUT / '博士毕业答辩_20261007.pptx'
BASE = OUT / '博士毕业答辩_预览_20261007'
records = json.loads((HERE / 'slides.json').read_text(encoding='utf-8'))
prs = Presentation(PATH)
issues = []
for idx, (slide, record) in enumerate(zip(prs.slides, records), 1):
    if record['title'] not in '\n'.join(s.text for s in slide.shapes if s.has_text_frame):
        issues.append(f'{idx}: title lost')
    if record['notes'] not in slide.notes_slide.notes_text_frame.text:
        issues.append(f'{idx}: notes lost')
    for shape in slide.shapes:
        if shape.left < 0 or shape.top < 0 or shape.left+shape.width > prs.slide_width+20 or shape.top+shape.height > prs.slide_height+20:
            issues.append(f'{idx}: off-slide shape {shape.name}')
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    if not cell.text:
                        issues.append(f'{idx}: empty table cell')
with zipfile.ZipFile(PATH) as z:
    slide_parts = [name for name in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', name)]
    native_math = sum(len(etree.fromstring(z.read(name)).xpath('//*[local-name()="oMath"]')) for name in slide_parts)
    charts = len([name for name in z.namelist() if re.fullmatch(r'ppt/charts/chart\d+\.xml', name)])
    math_intended = sum(len(r['equations']) for r in records)
    if native_math != math_intended:
        issues.append(f'Native equations {native_math} != intended {math_intended}')
with fitz.open(str(BASE)+'.pdf') as pdf:
    if len(pdf) != len(records):
        issues.append('PDF page count differs from deck')
    missing = []
    for idx, (p, r) in enumerate(zip(pdf, records), 1):
        text = p.get_text()
        (HERE / f'page_{idx:02d}.txt').write_text(text, encoding='utf-8')
        clean = lambda value: re.sub(r'\s+', '', value)
        if clean(r['title']) not in clean(text):
            missing.append(idx)
    if missing:
        issues.append(f'PDF titles not extracted: {missing}')

font = ImageFont.truetype('/mnt/c/Windows/Fonts/arial.ttf', 20)
imgs = sorted(Path(str(BASE)+'_pages').glob('slide_*.png'))
for batch in range(0, len(imgs), 10):
    sheet = Image.new('RGB', (1640, 2450), 'white')
    draw = ImageDraw.Draw(sheet)
    for j, path in enumerate(imgs[batch:batch+10]):
        with Image.open(path) as im:
            thumb = im.resize((800, 450), Image.Resampling.LANCZOS)
        x, y = 10+(j%2)*820, 35+(j//2)*485
        sheet.paste(thumb, (x, y))
        draw.text((x+4, y-26), f'{batch+j+1:02d}', font=font, fill='black')
    sheet.save(HERE / f'contact_{batch//10+1:02d}.jpg', quality=93)

pagination = json.loads(Path(str(BASE)+'_pagination.json').read_text(encoding='utf-8-sig'))
result = {
    'slides': len(prs.slides), 'main_slides': sum(not r['backup'] for r in records),
    'backup_slides': sum(r['backup'] for r in records), 'main_seconds': sum(r['seconds'] for r in records),
    'native_equations': native_math, 'native_charts': charts,
    'native_tables': sum(s.has_table for slide in prs.slides for s in slide.shapes),
    'notes': len(records), 'software': pagination['software'], 'version': pagination['version'],
    'build': pagination['build'], 'text_overflow': pagination['textOverflow'],
    'structure_issues': issues,
    'fonts_present': {name: Path('/mnt/c/Windows/Fonts', name).is_file() for name in ['msyh.ttc', 'cambria.ttc']},
}
(HERE / 'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False, indent=2))
