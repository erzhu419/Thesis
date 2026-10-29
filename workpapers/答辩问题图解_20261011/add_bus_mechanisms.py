"""Add the requested bus examples without rewriting the other sections."""
from copy import deepcopy
import json
from pathlib import Path
from shutil import copy2
import sys

from lxml import etree
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / 'workpapers/博士答辩_20261007'))
from drawing import Deck, TEAL, RED, PURPLE, GOLD, INK, MUTED, LINE, LIGHT, WHITE

SOURCE = ROOT / 'output/博士毕业答辩_问题图解_两图修订_20261011.pptx'
OUTPUT = ROOT / 'output/博士毕业答辩_公交原理补充_20261011.pptx'
Q_IMAGE = ROOT.parent / 'RE-SAC/paper/all_q_values_comparison_v2.png'
H2O_URL = 'https://arxiv.org/html/2309.12716#S4.SS2'


def new_page(title, section, source, notes):
    slide = deck.page(title, section, source, notes, 75)
    slide.shapes[4].text_frame.paragraphs[0].runs[0].hyperlink.address = H2O_URL if 'H2O+' in source else 'https://arxiv.org/abs/2603.18396'
    return slide


def q_failure():
    notes = (
        '左图直接使用RE-SAC论文中的公交Q估计图。仅认知分歧惩罚配置的估计发生深度下降，'
        '组合配置的中等集成保持较稳定。纵轴为对称对数尺度，不能按线性距离理解，也不是公交真实Q值的蒙特卡洛真值。'
        '右图解释一种可能的失配机制：相似状态与相同驻站动作仍可能遇到不同到达客流和道路延误；'
        '一次有利转移不代表该动作在所有运行中同样有效，后继价值与奖励噪声使更新目标波动。'
        '若额外的悲观惩罚过强，熟悉动作的估计可能被压得过低；陌生动作缺少回放约束，可能仍接近初始化尺度。'
        '在负成本奖励下，接近零的估计反而高于负的熟悉动作估计，使演员过度追逐未充分验证的动作。'
        '这不是神经网络必定把未见动作估为零的定理，图中-10、-80和零仅为机制示例。'
        '平稳随机环境中新一次采样可以不利，并不意味着环境规律已经改变；规律改变的问题留给BAPR。'
    )
    slide = new_page('随机反馈下，Q估计为何可能放弃已有的有效动作？',
        '第3章 | 平稳随机环境下的共享鲁棒驻站控制',
        'RE-SAC论文：公交Q估计原图；右侧数值仅用于解释机制', notes)
    deck.text(.8, 1.77, 6.5, .42, '公交训练中确实出现价值估计深度下降', 21, True, TEAL)
    deck.picture(Q_IMAGE, .73, 2.25, 6.95, 3.82)
    deck.line(7.87, 1.88, 7.87, 5.92, LINE, 1)
    deck.text(8.12, 1.77, 4.45, .45, '同一驻站动作，运行后果有波动', 20, True, PURPLE)
    deck.bus(8.24, 2.49, .55, color=TEAL)
    deck.text(9.02, 2.47, 3.33, .41, '相似状态 s，驻站动作 a', 17)
    deck.node(8.12, 3.11, 1.96, 1.01, '曾经较好', '奖励示例 -10', TEAL, 18)
    deck.node(10.54, 3.11, 1.96, 1.01, '本次较差', '奖励示例 -80', RED, 18)
    deck.line(9.09, 2.97, 9.09, 3.07, TEAL, 1.7, arrow=True)
    deck.line(9.5, 2.97, 11.5, 3.07, RED, 1.7, arrow=True)
    deck.text(8.18, 4.19, 4.37, .42, '噪声 + 过强悲观惩罚', 20, True, RED, PP_ALIGN.CENTER)
    deck.line(10.24, 4.7, 10.24, 4.89, RED, 1.8, arrow=True)
    deck.text(8.13, 4.99, 4.49, .49, '若熟悉动作被压低，陌生动作≈0', 17, True)
    deck.text(8.17, 5.62, 4.38, .4, '负奖励下，陌生动作反而显得更好', 17, True, PURPLE, PP_ALIGN.CENTER)
    deck.verdict('RE-SAC要保护已有经验的利用，同时限制少见动作的不可靠价值估计。')
    return slide


def h2o_ratio():
    notes = (
        '本页解释Niu等H2O+原方法式(8)和式(9)，不是新增公交结果。'
        '权重的分子是目标环境中、给定当前状态和动作后出现该后继结果的概率密度；分母是源仿真的对应概率密度。'
        '因此判断的是条件转移规律，不是把状态向量与离线数据计算外观相似度。'
        '同一时间窗到站的两结果示例中，目标到站概率0.8、源仿真0.4，权重为2；'
        '目标未到站概率0.2、仿真0.6，权重为三分之一。它会增强目标支持而源端采得不足的结果，减弱源端过量的结果。'
        '加权不是天然的硬删除；只有另设阈值或权重为零时才等价于筛除。'
        '原框架还混合离线状态价值与SAC后继价值，并非只有这一比率。'
        '本文公交正式完整控制器使用单位源权重，分类器比率保留为诊断；不能用这张原方法原理页解释未启用项的实验收益。'
    )
    slide = new_page('H2O+原理：目标更支持的仿真后果，学习权重更高',
        '第5章 | 固定目标数据支持的跨保真迁移',
        'Niu等H2O+式(8)–(9)；原方法原理，本文公交主结果采用单位源权重', notes)
    deck.text(.89, 1.8, 11.5, .39, '给定相同的当前公交状态、驻站动作，比较同一种后继结果', 21, True, TEAL)
    deck.text(.91, 2.55, 3.15, .56, '仿真经验权重 =', 24, True, INK)
    deck.text(4.19, 2.33, 8.0, .48, '目标记录中，这种后继变化的概率（密度）', 23, True, PURPLE, PP_ALIGN.CENTER)
    deck.line(4.19, 2.97, 12.18, 2.97, INK, 1.7)
    deck.text(4.19, 3.15, 8.0, .48, '源仿真中，这种后继变化的概率（密度）', 23, True, RED, PP_ALIGN.CENTER)
    deck.text(.88, 4.08, 3.23, .4, '示例：一段时间内是否到站', 17, True, MUTED)
    deck.bus(1.12, 4.94, .87, color=RED)
    deck.text(.88, 5.68, 2.75, .3, '同一源仿真生成两类结果', 15, color=MUTED)
    for y, name, probability, result, color in [
        (4.1, '已经到站', '目标0.8 / 仿真0.4', '2：增强学习', TEAL),
        (5.17, '尚未到站', '目标0.2 / 仿真0.6', '1/3：减弱学习', RED),
    ]:
        deck.line(2.5, 5.09, 4.03, y+.35, LINE, 1.8, arrow=True)
        deck.text(4.17, y, 2.47, .4, name, 21, True, color)
        deck.text(6.38, y, 3.03, .4, probability, 18, color=INK)
        deck.line(9.46, y+.23, 9.87, y+.23, color, 2, arrow=True)
        deck.text(10.04, y-.01, 2.38, .67, result, 20, True, color, PP_ALIGN.CENTER)
    deck.verdict('源仿真提供可交互经验；目标记录帮助判断哪些运行后果更值得相信。')
    return slide


def arrival_row(x, y, label, color, arrived):
    deck.text(x, y, 3.45, .33, label, 17, True, color)
    road_y = y + .87
    stop_x = x + 2.65
    deck.line(x+.09, road_y, x+3.5, road_y, MUTED, 1.4)
    deck.line(stop_x, y+.43, stop_x, road_y, GOLD, 2)
    deck.rect(stop_x-.13, y+.4, .27, .18, LIGHT, GOLD)
    bus_x = stop_x-.62 if arrived else x+.22
    deck.bus(bus_x, y+.55, .51, color=color)
    if not arrived:
        deck.line(bus_x+.78, y+.74, stop_x-.2, y+.74, RED, 1.9, arrow=True, dashed=True)


def mc_bus_example():
    slide = deck.prs.slides[31]
    deck.slide = slide
    for shape in list(slide.shapes)[7:]:
        shape._element.getparent().remove(shape._element)
    title = '目标已到站，仿真尚未到：修正预测，再判断可信程度'
    source = 'MC-WM论文：残差模型与可信权重；目标环境在本文中以SUMO代理验证'
    notes = (
        '沿用用户给出的三段式参考图，换成公交到站示例。比较同一状态与动作后的对应后继：'
        '目标记录中的公交已到站，源仿真预测仍在站前；残差标签是目标后继减去源预测。'
        '修正方向应是把未到的预测推向已到，而不是让算法把已到当成未到。'
        '原图的到站位置只作概念示意；公交实际模型预测事件、状态、时间与成本，不一定以该图的车辆位置作为输出坐标。'
        '修正从固定目标离线记录学习，不在每次源交互时读取目标环境的现场答案，也不改写SUMO实际轨迹。'
        '第二部分用留出目标记录检查修正预测的误差，形成单步置信并沿后继传播；可信权重是学习信号而不是已校准概率。'
        '第三部分把可信模型经验以较高权重加入评论家回归，低可信经验减弱影响，再由价值学习改善驻站策略。'
        '模型修正与模型使用强度是两个问题，后续保留归一化权重和公交服务结果页。'
    )
    slide.shapes[2].text_frame.paragraphs[0].runs[0].text = title
    slide.shapes[4].text_frame.paragraphs[0].runs[0].text = source
    for x in [4.82, 8.63]:
        deck.line(x, 1.89, x, 5.97, LINE, 1)
    deck.text(.88, 1.79, 3.74, .43, '1  用残差修正到站预测', 21, True, TEAL)
    arrival_row(.89, 2.38, '目标历史记录：已经到站', TEAL, True)
    arrival_row(.89, 3.47, '源仿真预测：尚未到站', RED, False)
    arrival_row(.89, 4.61, '源预测 + 残差：已到站', PURPLE, True)
    deck.text(.92, 5.66, 3.5, .36, '补上仿真遗漏的后继变化', 17, True, PURPLE)

    deck.text(5.03, 1.79, 3.47, .43, '2  核对误差，学习可信度', 21, True, PURPLE)
    deck.text(5.13, 2.42, 3.28, .96, '修正预测\n对照留出目标记录', 22, True, INK, PP_ALIGN.CENTER)
    deck.line(6.8, 3.46, 6.8, 3.67, PURPLE, 1.8, arrow=True)
    deck.text(5.16, 3.78, 3.18, .45, '误差小 → 单步置信高', 19, True, TEAL, PP_ALIGN.CENTER)
    deck.text(5.16, 4.44, 3.18, .45, '误差大 → 单步置信低', 19, True, RED, PP_ALIGN.CENTER)
    for j, color in enumerate([TEAL, GOLD, TEAL]):
        x = 5.29 + j*1.15
        deck.ellipse(x, 5.24, .28, .28, color)
        if j < 2:
            deck.line(x+.33, 5.38, x+.94, 5.38, MUTED, 1.7, arrow=True)
    deck.text(5.11, 5.74, 3.21, .33, '结合后继的可信度继续判断', 16, color=MUTED, align=PP_ALIGN.CENTER)

    deck.text(8.91, 1.79, 3.52, .43, '3  加权参与策略训练', 21, True, GOLD)
    deck.database(10.04, 2.49, 1.36, .69, GOLD)
    deck.text(9.04, 3.4, 3.2, .36, '模型生成补充转移', 20, True, GOLD, PP_ALIGN.CENTER)
    deck.text(8.97, 4.12, 3.43, .36, '可靠预测：提高学习权重', 19, True, TEAL)
    deck.line(9.17, 4.72, 12.02, 4.72, TEAL, 4.5, arrow=True)
    deck.text(8.97, 5.01, 3.43, .36, '不可靠预测：减弱影响', 19, True, RED)
    deck.line(9.17, 5.55, 12.02, 5.55, RED, 1.1, arrow=True)
    deck.text(8.96, 5.79, 3.37, .29, '评论家学习 → 驻站策略', 17, True, INK, PP_ALIGN.CENTER)
    deck.verdict('先让预测更接近目标运行规律，再让可靠的模型经验真正帮助策略学习。')
    return slide, title, source, notes


def without_footer(slide):
    xml = deepcopy(slide._element)
    for shape in xml.xpath('./p:cSld/p:spTree/p:sp'):
        if shape.xpath('./p:nvSpPr/p:cNvPr[@id="7"]'):
            shape.getparent().remove(shape)
    return etree.tostring(xml)


def main():
    backup = HERE / 'before_bus_mechanisms'
    backup.mkdir(exist_ok=True)
    if not (backup / SOURCE.name).exists():
        copy2(SOURCE, backup / SOURCE.name)
    global deck
    deck = Deck()
    deck.prs = Presentation(SOURCE)
    before = Presentation(SOURCE)
    original = list(deck.prs.slides)
    base_records = json.loads((HERE / 'slides.json').read_text())
    meta = {}
    for i, slide in enumerate(original, 1):
        record = deepcopy(base_records[i-1])
        record['title'] = slide.shapes[2].text
        record['source'] = slide.shapes[4].text
        note_lines = slide.notes_slide.notes_text_frame.text.splitlines()
        record['notes'] = '\n'.join(note_lines[1:-1])
        record['previous_page'] = i
        meta[id(slide)] = record
    q_slide = q_failure()
    meta[id(q_slide)] = deepcopy(deck.records[-1]) | {'previous_page':'new_Q'}
    h2o_slide = h2o_ratio()
    meta[id(h2o_slide)] = deepcopy(deck.records[-1]) | {'previous_page':'new_H2O_ratio'}
    slide, title, source, notes = mc_bus_example()
    meta[id(slide)].update(title=title, source=source, notes=notes)
    order = []
    for i, slide in enumerate(original, 1):
        order.append(slide)
        if i == 11:
            order.append(q_slide)
        if i == 26:
            order.append(h2o_slide)
    identifiers = {item.rId: item for item in list(deck.prs.slides._sldIdLst)}
    for item in list(deck.prs.slides._sldIdLst):
        deck.prs.slides._sldIdLst.remove(item)
    for slide in order:
        rid = next(rel.rId for rel in deck.prs.part.rels.values() if not rel.is_external and rel.target_part is slide.part)
        deck.prs.slides._sldIdLst.append(identifiers[rid])

    main_records = [meta[id(slide)] for slide in order if not meta[id(slide)].get('backup')]
    total = sum(r['seconds'] for r in main_records)
    times = [round(r['seconds'] * 2400 / total) for r in main_records]
    times[-2] += 2400 - sum(times)
    elapsed = 0
    records = []
    for i, slide in enumerate(order, 1):
        record = meta[id(slide)]
        record['page'] = i
        slide.shapes[5].text_frame.paragraphs[0].runs[0].text = f'{i:02d}'
        if record.get('backup'):
            prefix = '备用页，不计入主讲。'
            record['seconds'] = 0
        else:
            seconds = times[len(records)]
            record.update(seconds=seconds, start_seconds=elapsed, end_seconds=elapsed+seconds)
            elapsed += seconds
            prefix = f'建议讲述时间：{seconds}秒。'
        slide.notes_slide.notes_text_frame.text = prefix + '\n' + record['notes'] + '\n来源：' + record['source']
        if slide in [q_slide, h2o_slide, original[31]]:
            for effect in slide._element.xpath('.//a:effectRef'):
                effect.set('idx', '0')
        records.append(record)
    deck.prs.save(OUTPUT)
    (HERE / 'bus_mechanism_slides.json').write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding='utf-8')
    lines = ['# 博士毕业答辩讲稿：公交原理补充', '', '主讲45页，约40分钟；12页备用。', '']
    for record in records:
        if record.get('backup'):
            stamp = '备用'
        else:
            a, b = record['start_seconds'], record['end_seconds']
            stamp = f'{a//60:02d}:{a%60:02d}–{b//60:02d}:{b%60:02d}'
        lines += [f"## {record['page']:02d} {record['title']}（{stamp}）", '', record['notes'], '', '来源：'+record['source'], '']
    (ROOT / 'output/博士毕业答辩_公交原理补充_讲稿_20261011.md').write_text('\n'.join(lines), encoding='utf-8')

    # Reject unintended edits to the existing diagram/result objects.
    reopened = Presentation(OUTPUT)
    index = {r['previous_page']:r['page'] for r in records}
    unchanged = all(without_footer(before.slides[i-1]) == without_footer(reopened.slides[index[i]-1]) for i in range(1,56) if i != 32)
    assert unchanged and len(reopened.slides) == 57 and elapsed == 2400
    assert SOURCE.read_bytes() == (backup / SOURCE.name).read_bytes()
    cs_charts = [(i, s) for i in range(18,25) for s in before.slides[i-1].shapes if s.has_chart]
    for page, shape in cs_charts:
        counterpart = next(s for s in reopened.slides[index[page]-1].shapes if s.shape_id == shape.shape_id)
        assert etree.tostring(shape.chart._element) == etree.tostring(counterpart.chart._element)
    result = {'slides':57, 'main_slides':45, 'main_seconds':2400,
              'changed_pages':[index['new_Q'], index['new_H2O_ratio'], index[32]],
              'other_existing_slide_objects_unchanged':unchanged,
              'CS_BAPR_charts_unchanged':True, 'source_preserved':True}
    (HERE / 'bus_mechanism_verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
