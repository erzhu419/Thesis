from pathlib import Path
from copy import deepcopy
import json

from lxml import etree
from latex2mathml.converter import convert
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE.parents[1] / 'output'
INK = '263238'
MUTED = '65747A'
TEAL = '087F8C'
RED = 'C9654B'
PURPLE = '78618C'
GOLD = 'B08A36'
GRAY = 'A7B2B6'
LINE = 'DCE3E5'
LIGHT = 'EFF5F6'
WHITE = 'FFFFFF'
FONT = 'Microsoft YaHei'
W, H = 13.333333, 7.5
MATH_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
DRAW_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
WORD_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
MATH_TRANSFORM = etree.XSLT(etree.parse('/mnt/c/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'))


def rgb(value):
    return RGBColor.from_string(value)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.prs.core_properties.title = '从仿真到现实的公交系统驻站控制方法研究'
        self.prs.core_properties.author = '张异凡'
        self.prs.core_properties.subject = '中南大学博士学位论文答辩'
        self.records = []
        self.slide = None

    def page(self, title, section, source='', notes='', seconds=55, backup=False):
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = rgb(WHITE)
        i = len(self.prs.slides)
        self.rect(.5, .32, .065, .23, TEAL, None)
        self.text(.72, .23, 10.7, .35, section, 11, color=MUTED)
        self.text(.6, .78, 12.1, .62, title, 28, bold=True)
        self.line(.6, 1.53, 12.7, 1.53, LINE, 1)
        self.text(.6, 7.03, 11.5, .24, source, 9, color=MUTED)
        self.text(12.05, 7.0, .65, .3, f'{i:02d}', 11, color=MUTED, align=PP_ALIGN.RIGHT)
        self.line(.6, 6.93, 12.7, 6.93, LINE, .7)
        self.slide.notes_slide.notes_text_frame.text = (
            ('备用页，不计入主讲。\n' if backup else f'建议讲述时间：{seconds}秒。\n')
            + notes + '\n来源：' + source
        )
        self.records.append({'page': i, 'title': title, 'section': section, 'source': source,
                             'notes': notes, 'seconds': 0 if backup else seconds, 'backup': backup,
                             'equations': []})
        return self.slide

    def text(self, x, y, w, h, value, size=20, bold=False, color=INK,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, italic=False):
        shape = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = Inches(.015)
        tf.margin_top = tf.margin_bottom = Inches(.015)
        tf.vertical_anchor = valign
        for i, line in enumerate(str(value).split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_before = Pt(0)
            p.space_after = Pt(0)
            p.line_spacing = 1.15
            r = p.add_run()
            r.text = line
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = rgb(color)
            ea = OxmlElement('a:ea')
            ea.set('typeface', FONT)
            r._r.get_or_add_rPr().append(ea)
        return shape

    def rect(self, x, y, w, h, fill=None, line=LINE, radius=False):
        typ = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        s = self.slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill:
            s.fill.solid()
            s.fill.fore_color.rgb = rgb(fill)
        else:
            s.fill.background()
        if line:
            s.line.color.rgb = rgb(line)
            s.line.width = Pt(1.1)
        else:
            s.line.fill.background()
        if radius:
            s.adjustments[0] = .08
        return s

    def ellipse(self, x, y, w, h, fill=TEAL, line=None):
        s = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        if line:
            s.line.color.rgb = rgb(line)
        else:
            s.line.fill.background()
        return s

    def line(self, x1, y1, x2, y2, color=INK, width=1.7, arrow=False, dashed=False):
        s = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        s.line.color.rgb = rgb(color)
        s.line.width = Pt(width)
        ln = s._element.spPr.find('{%s}ln' % DRAW_NS)
        if arrow:
            end = OxmlElement('a:tailEnd')
            end.set('type', 'triangle')
            ln.append(end)
        if dashed:
            dash = OxmlElement('a:prstDash')
            dash.set('val', 'dash')
            ln.append(dash)
        return s

    def node(self, x, y, w, h, title, sub='', color=TEAL, size=20):
        self.rect(x, y, w, h, LIGHT, color, True)
        if sub:
            title_h = .4 if h <= 1.15 else .52
            sub_y = .57 if h <= 1.15 else .71
            self.text(x+.13, y+.09, w-.26, title_h, title, size, True, color, PP_ALIGN.CENTER)
            self.text(x+.13, y+sub_y, w-.26, h-sub_y-.08, sub, 16, color=INK, align=PP_ALIGN.CENTER)
        else:
            self.text(x+.12, y+.09, w-.24, h-.18, title, size, True, color, PP_ALIGN.CENTER)

    def equation(self, x, y, w, h, latex, size=25, color=INK):
        shape = self.text(x, y, w, h, '', size)
        p = shape.text_frame.paragraphs[0]
        p._p.clear()
        props = OxmlElement('a:pPr')
        props.set('algn', 'l')
        p._p.append(props)
        mathml = etree.fromstring(convert(latex).encode())
        result = MATH_TRANSFORM(mathml).getroot()
        container = etree.Element('{http://schemas.microsoft.com/office/drawing/2010/main}m',
                                  nsmap={'a14': 'http://schemas.microsoft.com/office/drawing/2010/main',
                                         'm': MATH_NS, 'a': DRAW_NS})
        if result.tag == '{%s}oMathPara' % MATH_NS:
            para = result
        else:
            para = etree.Element('{%s}oMathPara' % MATH_NS)
            para.append(result)
        jcprops = etree.Element('{%s}oMathParaPr' % MATH_NS)
        jc = etree.SubElement(jcprops, '{%s}jc' % MATH_NS)
        jc.set('{%s}val' % MATH_NS, 'left')
        para.insert(0, jcprops)
        for run in para.iter('{%s}r' % MATH_NS):
            for old in list(run):
                if old.tag == '{%s}rPr' % WORD_NS:
                    run.remove(old)
            rp = etree.Element('{%s}rPr' % DRAW_NS)
            rp.set('sz', str(int(size * 100)))
            rp.set('lang', 'en-US')
            solid = etree.SubElement(rp, '{%s}solidFill' % DRAW_NS)
            etree.SubElement(solid, '{%s}srgbClr' % DRAW_NS).set('val', color)
            etree.SubElement(rp, '{%s}latin' % DRAW_NS).set('typeface', 'Cambria Math')
            etree.SubElement(rp, '{%s}ea' % DRAW_NS).set('typeface', FONT)
            run.insert(0, rp)
        container.append(para)
        p._p.append(container)
        end = OxmlElement('a:endParaRPr')
        end.set('sz', str(int(size * 100)))
        p._p.append(end)
        self.records[-1]['equations'].append(latex)
        return shape

    def bus(self, x, y, scale=1, label='', color=TEAL):
        self.rect(x, y, 1.12*scale, .44*scale, color, None, True)
        for j in range(4):
            self.rect(x+(.12+j*.24)*scale, y+.07*scale, .17*scale, .15*scale, WHITE, None)
        self.ellipse(x+.15*scale, y+.37*scale, .15*scale, .15*scale, INK)
        self.ellipse(x+.81*scale, y+.37*scale, .15*scale, .15*scale, INK)
        if label:
            self.text(x-.2, y+.64*scale, 1.5*scale, .46, label, 16, color=color, align=PP_ALIGN.CENTER)

    def network(self, x, y, scale=1, color=TEAL):
        layers = [[(x, y+j*.4*scale) for j in range(3)],
                  [(x+.64*scale, y-.2*scale+j*.4*scale) for j in range(4)],
                  [(x+1.28*scale, y+j*.4*scale) for j in range(3)]]
        for left, right in zip(layers, layers[1:]):
            for a in left:
                for b in right:
                    self.line(a[0]+.08*scale, a[1]+.08*scale, b[0]+.08*scale, b[1]+.08*scale, LINE, .65)
        for layer in layers:
            for xx, yy in layer:
                self.ellipse(xx, yy, .16*scale, .16*scale, color)

    def database(self, x, y, w=1.1, h=.8, color=TEAL):
        self.rect(x, y+.13, w, h-.22, LIGHT, color)
        self.ellipse(x, y, w, .26, LIGHT, color)
        self.ellipse(x, y+h-.26, w, .26, LIGHT, color)
        self.line(x, y+.4, x+w, y+.4, color, 1)

    def verdict(self, value, color=TEAL):
        self.rect(.65, 6.22, 12.0, .48, LIGHT, None)
        self.rect(.65, 6.22, .05, .48, color, None)
        self.text(.83, 6.25, 11.6, .40, value, 19, bold=True, color=color)

    def chart(self, x, y, w, h, categories, series, kind='bar', colors=None,
              ymin=None, ymax=None, number_format='0', legend=True):
        data = CategoryChartData()
        data.categories = categories
        for name, values in series:
            data.add_series(name, values)
        typ = {'bar': XL_CHART_TYPE.COLUMN_CLUSTERED, 'line': XL_CHART_TYPE.LINE_MARKERS,
               'horizontal': XL_CHART_TYPE.BAR_CLUSTERED}[kind]
        chart = self.slide.shapes.add_chart(typ, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
        chart.has_legend = legend and len(series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.name = FONT
            chart.legend.font.size = Pt(13)
        chart.chart_style = 2
        chart.has_title = False
        chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        for axis in (chart.category_axis, chart.value_axis):
            axis.tick_labels.font.name = FONT
            axis.tick_labels.font.size = Pt(13)
            axis.tick_labels.font.color.rgb = rgb(MUTED)
            axis.format.line.color.rgb = rgb(LINE)
        chart.value_axis.tick_labels.number_format = number_format
        if ymin is not None:
            chart.value_axis.minimum_scale = ymin
        if ymax is not None:
            chart.value_axis.maximum_scale = ymax
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = rgb(LINE)
        palette = colors or [TEAL, RED, PURPLE, GOLD]
        for idx, s in enumerate(chart.series):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = rgb(palette[idx % len(palette)])
            s.format.line.color.rgb = rgb(palette[idx % len(palette)])
            if kind in ('bar', 'horizontal'):
                s.invert_if_negative = False
            if kind == 'line':
                s.format.line.width = Pt(2.5)
                s.marker.size = 6
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = rgb(palette[idx % len(palette)])
                s.marker.format.line.color.rgb = rgb(palette[idx % len(palette)])
        return chart

    def table(self, x, y, w, h, headers, rows, widths=None, size=16):
        shape = self.slide.shapes.add_table(len(rows)+1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
        table = shape.table
        if widths:
            for c, width in zip(table.columns, widths):
                c.width = Inches(width)
        for i, row in enumerate([headers] + rows):
            for j, value in enumerate(row):
                cell = table.cell(i, j)
                cell.text = str(value)
                cell.margin_left = Inches(.08)
                cell.margin_right = Inches(.08)
                cell.margin_top = Inches(.04)
                cell.margin_bottom = Inches(.04)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(TEAL if i == 0 else (LIGHT if i % 2 else WHITE))
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.LEFT
                    for r in p.runs:
                        r.font.name = FONT
                        r.font.size = Pt(size)
                        r.font.bold = i == 0
                        r.font.color.rgb = rgb(WHITE if i == 0 else INK)
        return shape

    def save(self, name='博士毕业答辩_20261007.pptx'):
        OUT.mkdir(exist_ok=True)
        path = OUT / name
        self.prs.save(path)
        (HERE / 'slides.json').write_text(json.dumps(self.records, ensure_ascii=False, indent=2), encoding='utf-8')
        print(path)
        return path
