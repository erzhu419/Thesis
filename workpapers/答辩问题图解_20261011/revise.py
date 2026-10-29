"""Scoped visual revision of the existing, editable defense deck."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import json
import math
from shutil import copy2
import sys

import fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
OLD = ROOT / 'workpapers/博士答辩_20261007'
sys.path.insert(0, str(OLD))
from drawing import Deck, TEAL, RED, PURPLE, GOLD, INK, MUTED, LINE, LIGHT, WHITE

SOURCE = ROOT / 'output/博士毕业答辩_批注修订_20261007.pptx'
OUTPUT = ROOT / 'output/博士毕业答辩_问题图解_20261011.pptx'
PAPER = ROOT.parent / 'Catetorical-Feature-SAC-in-bus-simulation/Single Agent Robust Deep Reinforcement Learning for Bus Fleet Control.pdf'
ASSETS = HERE / 'assets'
BRIGHT_RED = 'D72D3B'


def extract_trajectory():
    ASSETS.mkdir(exist_ok=True)
    with fitz.open(PAPER) as pdf:
        page = pdf[15]
        image = page.get_images()[0]
        payload = pdf.extract_image(image[0])['image']
    with Image.open(BytesIO(payload)) as im:
        # Both panels are embedded as one image. Keep (a), its axes and legend.
        panel = im.crop((0, 0, im.width, im.height // 2))
        panel.thumbnail((3300, 1650), Image.Resampling.LANCZOS)
        panel.save(ASSETS / 'categorical_figure6a.png')


def path(points, color, width=2, fill=None, close=False):
    builder = deck.slide.shapes.build_freeform(Inches(points[0][0]), Inches(points[0][1]))
    builder.add_line_segments([(Inches(x), Inches(y)) for x, y in points[1:]], close=close)
    shape = builder.convert_to_shape()
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.fill.transparency = .7
    else:
        shape.fill.background()
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width)
    return shape


def bell(x, y, w, h, mean=0, sigma=.7, color=TEAL, bounds=(-3, 3), peak=1, baseline=True):
    low, high = bounds
    points = []
    for i in range(81):
        xx = low + (high - low) * i / 80
        value = math.exp(-.5 * ((xx - mean) / sigma) ** 2) / (sigma * math.sqrt(2 * math.pi))
        points.append((x + w * i / 80, y + h - h * value / peak))
    if baseline:
        deck.line(x, y + h, x + w, y + h, LINE, .8)
    return path(points, color, 2.5)


def person(x, y, color=RED, scale=1):
    deck.ellipse(x + .065 * scale, y, .12 * scale, .12 * scale, color)
    deck.line(x + .125 * scale, y + .14 * scale, x + .125 * scale, y + .37 * scale, color, 2.2)
    deck.line(x + .125 * scale, y + .2 * scale, x + .035 * scale, y + .29 * scale, color, 2.2)
    deck.line(x + .125 * scale, y + .2 * scale, x + .22 * scale, y + .29 * scale, color, 2.2)
    deck.line(x + .125 * scale, y + .37 * scale, x + .035 * scale, y + .52 * scale, color, 2.2)
    deck.line(x + .125 * scale, y + .37 * scale, x + .22 * scale, y + .52 * scale, color, 2.2)


def span(x1, x2, y, label, color=TEAL, size=17):
    deck.line(x1, y, x2, y, color, 2.2, arrow=True)
    deck.line(x1, y - .06, x1, y + .06, color, 1.4)
    deck.line(x2, y - .06, x2, y + .06, color, 1.4)
    deck.text(x1, y - .42, x2 - x1, .32, label, size, True, color, PP_ALIGN.CENTER)


def use(key, title=None, source=None, notes=None):
    slide = slides[key]
    deck.slide = slide
    deck.records = [records[key]]
    for shape in list(slide.shapes)[7:]:
        shape._element.getparent().remove(shape._element)
    record = records[key]
    if title:
        record['title'] = title
        slide.shapes[2].text_frame.paragraphs[0].runs[0].text = title
    if source:
        record['source'] = source
        slide.shapes[4].text_frame.paragraphs[0].runs[0].text = source
    if notes:
        record['notes'] = notes
    record['equations'] = []
    return slide


def reward_slide():
    use(5, notes='左图沿用第2页同站比较相邻车辆的思路：前车、本车、后车的服务时刻形成前后时距，计划时距目标要求两者接近计划间隔，避免串车和过大的服务空档。右图把乘客时间分成站内候车与车内时间：候车累计按两倍权重，车内累计按一倍权重。车辆额外驻站可能缩小下游服务空档，也会延长当前车内乘客时间。两类奖励各回答不同运营问题，所以后续展示同一评价协议内的结果，不把奖励变好等同于所有指标都变好。图中乘客与间隔仅解释定义，不是新增实验。')
    deck.text(.8, 1.78, 5.8, .46, '计划时距：让相邻服务间隔更均衡', 22, True, TEAL)
    deck.text(7.0, 1.78, 5.4, .46, '乘客时间：减少候车与车内耗时', 22, True, RED)
    deck.line(6.65, 1.83, 6.65, 5.88, LINE, 1)
    deck.line(1.1, 3.78, 6.08, 3.78, MUTED, 1.5, arrow=True)
    for x, label, color in [(1.2, '前车服务', TEAL), (3.22, '本车服务', GOLD), (5.24, '后车服务', PURPLE)]:
        deck.bus(x - .33, 2.59, .68, color=color)
        deck.line(x + .06, 3.16, x + .06, 3.88, color, 1.1, dashed=True)
        deck.text(x - .68, 4.02, 1.5, .36, label, 17, True, color, PP_ALIGN.CENTER)
    span(1.26, 3.28, 3.51, '前向时距 hᶠ')
    span(3.28, 5.3, 3.51, '后向时距 hᵇ', PURPLE)
    deck.text(5.25, 4.42, .9, .3, '时间 →', 14, color=MUTED)
    deck.equation(1.0, 4.75, 5.2, .62, r'h^f\approx h^b\approx h_0', 27)
    deck.text(.93, 5.42, 5.48, .48, 'h₀：计划间隔；偏离目标与前后失衡都受惩罚', 17, color=MUTED)
    # One continuous passenger journey: platform queue, boarding, then ride.
    deck.rect(7.08, 2.55, .64, .53, LIGHT, TEAL)
    deck.text(7.08, 2.55, .64, .53, '站点', 15, True, TEAL, PP_ALIGN.CENTER)
    for j in range(4):
        person(7.85 + .4*j, 2.48, RED, .92)
    deck.line(7.03, 3.12, 9.62, 3.12, TEAL, 2)
    deck.bus(10.05, 2.43, 1.75, color=TEAL)
    for j in range(4):
        deck.ellipse(10.37 + j*.38, 2.73, .11, .11, WHITE)
    deck.line(9.48, 2.91, 9.95, 2.91, RED, 2, arrow=True)
    deck.text(7.28, 3.51, 2.2, .36, '站内排队候车', 19, True, RED, PP_ALIGN.CENTER)
    deck.text(10.02, 3.51, 2.0, .36, '车内乘坐', 19, True, TEAL, PP_ALIGN.CENTER)
    deck.rect(7.18, 4.0, 2.15, .56, 'FBECE8', None)
    deck.text(7.18, 4.0, 2.15, .56, '候车时间 × 2', 20, True, RED, PP_ALIGN.CENTER)
    deck.text(9.4, 4.02, .42, .5, '+', 23, True, MUTED, PP_ALIGN.CENTER)
    deck.rect(9.91, 4.0, 2.4, .56, LIGHT, None)
    deck.text(9.91, 4.0, 2.4, .56, '车内时间 × 1', 20, True, TEAL, PP_ALIGN.CENTER)
    deck.equation(7.06, 4.83, 5.4, .56, r'G=2W+T^{\mathrm{ride}}', 27)
    deck.text(7.06, 5.42, 5.35, .5, '训练奖励取负的累计成本；评价报告人均成本', 17, color=MUTED)
    deck.verdict('驻站既能改善后续间隔，也会增加车内时间：两种目标必须分别评价。')


def variance_slide():
    oldslide = deck.slide
    section = records[7]['section']
    source = '前序公交论文 Figure 6(a)原图；下方分布为原理示意，不是实测方差'
    notes = ('上方为前序公交论文Figure 6(a)的单次受控运行轨迹，不是重复试验的方差图。'
             '8点与18点的引线用于定位运营时段；比较的是同一站点、相同服务位置在各次重复运行中的状态波动，并非一条轨迹里同一个班次同时出现在两个时段。'
             '中间用早期较窄、后期较宽的分布说明行驶时间与客流随机扰动可能沿运行过程累积；方差增大并非任何控制器下都严格单调，图中没有新增方差数值。'
             '下方说明：在宽分布上汇总训练，有限数据覆盖和缺少模式信息可能让一个平均策略在某个部署条件下失配。'
             '两个子分布表示不同条件下的访问区域，而不是声称从同一分布独立抽一批测试样本就必然退化。'
             '红叉标示需要防范的失败情形；RE-SAC用共享经验和保守评价降低对局部高价值估计的依赖。')
    slide = deck.page('随机扰动会累积，训练与部署可能落在不同区域', section, source, notes, 80)
    records['variance'] = deck.records[-1]
    slides['variance'] = slide
    deck.text(.79, 1.73, 11.9, .29, '受控公交轨迹：从早间到晚间的完整运营过程', 17, True, TEAL)
    deck.picture(ASSETS / 'categorical_figure6a.png', 1.04, 1.94, 11.34, 2.12)
    # Arrow locations correspond to the hourly ticks in the original panel.
    deck.line(4.23, 4.02, 3.42, 4.24, TEAL, 1.8, arrow=True)
    deck.line(9.86, 4.02, 9.43, 4.24, RED, 1.8, arrow=True)
    deck.text(1.15, 4.19, 4.5, .34, '08:00：较窄的状态分布', 19, True, TEAL, PP_ALIGN.CENTER)
    deck.text(7.02, 4.19, 4.5, .34, '18:00：可能累积更大波动', 19, True, RED, PP_ALIGN.CENTER)
    bell(1.2, 4.56, 4.42, .6, sigma=.52, color=TEAL, peak=.9)
    bell(7.1, 4.56, 4.42, .6, sigma=1.42, color=RED, peak=.9)
    deck.text(5.72, 4.52, 1.24, .61, '扰动累积\n→', 16, True, MUTED, PP_ALIGN.CENTER)
    deck.text(.83, 5.19, 11.62, .25, '同一站点、相同服务位置的重复运行分布；曲线仅示意宽度变化', 12, color=MUTED, align=PP_ALIGN.CENTER)
    bell(.87, 5.46, 2.33, .34, sigma=1.42, color=LINE, peak=1)
    bell(.87, 5.46, 2.33, .34, mean=-1.05, sigma=.44, color=TEAL, peak=1, baseline=False)
    deck.text(.92, 5.83, 2.22, .3, '训练访问区域', 16, True, TEAL, PP_ALIGN.CENTER)
    deck.line(3.23, 5.83, 3.8, 5.83, INK, 1.8, arrow=True)
    deck.text(3.83, 5.59, 2.35, .47, '共享策略 π', 20, True, INK, PP_ALIGN.CENTER)
    deck.line(6.2, 5.83, 6.85, 5.83, INK, 1.8, arrow=True)
    bell(6.87, 5.46, 2.6, .34, sigma=1.42, color=LINE, peak=1)
    bell(6.87, 5.46, 2.6, .34, mean=1.05, sigma=.44, color=RED, peak=1, baseline=False)
    deck.text(6.93, 5.83, 2.46, .3, '不同部署访问区域', 16, True, RED, PP_ALIGN.CENTER)
    deck.text(9.65, 5.57, 2.21, .5, '可能失配', 23, True, RED, PP_ALIGN.CENTER)
    deck.line(12.02, 5.61, 12.43, 6.02, BRIGHT_RED, 5)
    deck.line(12.02, 6.02, 12.43, 5.61, BRIGHT_RED, 5)
    deck.verdict('状态覆盖很宽，并不意味着一个策略在每种局部运行条件下都可靠。')
    deck.slide = oldslide


def resac_problem():
    use(7, 'RE-SAC：客流与路况随机，动作后果并不固定', notes='先讲第三章面对的公交问题。即使给出同样的驻站时长，不同车辆、站点和运营时段仍会遇到不同的客流与路况。随机扰动可能沿运营过程累积，有限经验中的动作价值也可能过度乐观。因此先让车辆共享经验并保留类别身份，再用鲁棒集成软演员评论家RE-SAC的集成与保守评价降低对不可靠价值的依赖。下一页用轨迹和分布示意说明为什么覆盖和后果波动是问题，后面用结果回答方法是否改善服务。')
    deck.text(.86, 1.93, 11.6, .42, '同样额外驻站30秒，后续服务为什么不同？', 24, True, TEAL)
    deck.bus(1.38, 3.1, 1.0, color=TEAL)
    deck.text(.85, 4.01, 2.65, .43, '相同驻站动作', 21, True, TEAL, PP_ALIGN.CENTER)
    for y, name, sub, color in [(2.81, '客流到达随机', '上车人数与停站服务变化', RED),
                                 (3.94, '道路运行随机', '同一路段也有不同延误', PURPLE),
                                 (5.07, '经验覆盖不均', '少见状态的价值更不可靠', GOLD)]:
        deck.line(2.85, 3.63, 4.25, y+.39, LINE, 1.7, arrow=True)
        deck.text(4.48, y, 3.66, .52, name, 23, True, color)
        deck.text(8.12, y, 4.1, .43, sub, 19)
    deck.verdict('RE-SAC：保留车辆与站点身份，共享经验，并保守评价不确定的后果。')


def resac_mechanism():
    use(11, '面对噪声与少见状态，怎样避免过度乐观？', notes='RE-SAC让多个价值网络分别评价同一驻站候选。意见接近时，经验对动作后果的支持较一致；意见差异大时，不能只追求其中最高的估计。方法在后继价值更新与当前动作评价两个位置加入保守调节，同时保留成员自己的目标估计。这里的集成分歧和参数范数是计算代理，不直接等于前一页贝叶斯方法分离出的两种真实方差。详细目标和两个系数放在备用页，主讲只解释该机制怎样帮助公交避免依赖不可靠的动作价值。')
    deck.node(.85, 2.56, 2.72, 1.32, '当前公交状态', '比较候选驻站时长', TEAL)
    for y, label in [(1.95, '价值网络1'), (3.19, '价值网络2'), (4.43, '价值网络K')]:
        deck.node(4.27, y, 2.53, .9, label, color=PURPLE, size=21)
        deck.line(3.57, 3.22, 4.19, y+.45, LINE, 1.6, arrow=True)
        deck.line(6.8, y+.45, 7.49, 3.26, LINE, 1.6, arrow=True)
    deck.node(7.55, 2.54, 2.38, 1.42, '合并与保守评价', '高分歧动作不盲信', GOLD, 18)
    deck.node(10.53, 2.54, 2.0, 1.42, '驻站决策', '收益与风险兼顾', TEAL, 19)
    deck.line(9.93, 3.26, 10.47, 3.26, INK, 1.8, arrow=True)
    deck.text(.99, 5.53, 11.32, .43, '共享经验提升覆盖；集成保留不同估计；保守调节限制不可靠的乐观判断。', 20, color=MUTED)
    deck.verdict('不是选“最高分”，而是在可用经验支持下选择可靠的驻站动作。')


def bapr_problem():
    use(15, 'BAPR：运行规律变了，先识别当前模式',
        '论文§4.1–§4.3；BAPR模式识别思想；分布为原理示意',
        '第三章处理给定规律中的随机波动；第四章进一步考虑规律本身改变。未识别模式时，不同条件的后果混在一起，使总分布显得很宽。引入模式变量z并从历史转移更新判断，可以区分模式内随机性和模式未知带来的混杂。这里画的是带标签的条件分布，不是等宽切割，也不保证每个条件方差都严格小于总方差。可理解为识别减少了可消除的模式不确定性，不是把物理偶然噪声变成认知不确定性。最新版BAPR在已知执行变换上验证：模式判断选择同一个参考策略的动作坐标补偿，不训练或切换n个互不相关的策略。公交统计反馈与该受控识别验证分别解释。')
    deck.text(.88, 1.81, 4.55, .43, '未识别：各种条件混在一起', 23, True, RED)
    deck.text(7.03, 1.81, 5.45, .43, '识别后：按模式理解后果', 23, True, PURPLE)
    # Mixture width comes both from within-mode scatter and distinct means.
    x, y, w, h = 1.03, 2.65, 4.44, 1.19
    means = [-1.4, 0, 1.4]
    values = []
    for i in range(121):
        xx = -3 + 6*i/120
        density = sum(math.exp(-.5*((xx-m)/.64)**2)/(.64*math.sqrt(2*math.pi))/3 for m in means)
        values.append((x+w*i/120, y+h-h*density/.36))
    deck.line(x, y+h, x+w, y+h, LINE, 1)
    path(values, RED, 3)
    deck.text(.93, 4.02, 4.68, .42, '宽分布：模式差异 + 模式内波动', 18, True, RED, PP_ALIGN.CENTER)
    deck.line(5.74, 3.37, 6.65, 3.37, PURPLE, 2.5, arrow=True)
    for j, (mean, color) in enumerate(zip(means, [TEAL, GOLD, PURPLE])):
        xx = 6.98+j*1.82
        deck.text(xx, 2.34, 1.65, .43, f'模式{j+1}', 20, True, color, PP_ALIGN.CENTER)
        bell(xx, 2.9, 1.65, .94, mean=mean, sigma=.64, color=color, peak=.8)
    deck.text(7.0, 4.02, 5.34, .42, '仍有随机波动，但不再把所有模式混为一谈', 17, color=MUTED, align=PP_ALIGN.CENTER)
    deck.node(.92, 5.01, 3.13, .79, '已经发生的转移', color=TEAL, size=20)
    deck.node(4.98, 5.01, 3.13, .79, '更新模式判断', color=PURPLE, size=20)
    deck.node(9.08, 5.01, 3.13, .79, '按当前模式响应', color=RED, size=20)
    deck.line(4.05, 5.4, 4.88, 5.4, INK, 1.8, arrow=True)
    deck.line(8.11, 5.4, 8.98, 5.4, INK, 1.8, arrow=True)
    deck.verdict('识别减少“模式未知”的认知不确定性；模式内的随机波动仍然存在。')


def bapr_identification():
    use(16, '不是给环境读编号，而是用运行后果辨认模式', notes='先看已完成的状态变化：同样发出一条指令，却出现怎样的变化？逆模型从前后状态推断名义执行信号，再与各候选模式对历史指令的解释比较，得到支持不同模式的证据。当前部署不会读取真实模式编号；编号只是候选关系的索引。模式的监督信息用于离线训练识别器，而不是在部署时直接提供答案。下一页保留原有滤波原理图：噪声观测不会立即改变判断，多次一致证据才让信念转向新模式。')
    deck.node(.85, 2.26, 3.08, 1.48, '已完成的状态变化', '前状态 → 后状态', TEAL)
    deck.node(5.12, 2.26, 3.12, 1.48, '推断实际执行效果', '这次变化由什么输入造成？', PURPLE, 18)
    deck.node(9.38, 2.26, 3.03, 1.48, '比较候选解释', '哪个模式最能解释？', RED)
    deck.line(3.93, 3.0, 5.0, 3.0, INK, 2, arrow=True)
    deck.line(8.24, 3.0, 9.25, 3.0, INK, 2, arrow=True)
    deck.text(1.05, 4.44, 3.04, .5, '历史指令也已知', 21, True, TEAL)
    deck.text(5.08, 4.42, 3.55, .76, '模式1、2、…、n\n分别给出候选执行效果', 20, color=PURPLE, align=PP_ALIGN.CENTER)
    deck.line(4.03, 4.7, 4.9, 4.7, TEAL, 1.8, arrow=True)
    deck.line(8.67, 4.69, 10.85, 3.89, RED, 1.8, arrow=True)
    deck.verdict('先用过去的转移形成证据，再决定下一次怎样响应，不能偷看真实模式。')


def residual_visual():
    use(36, '仿真低估了延误，目标记录怎样帮助纠偏？', notes='这一页讲源标定不足的具体问题，而不是先展示残差方程。示例中，源模型把下一事件间隔预测为300秒，目标记录显示为360秒，残差标签是60秒；数字只解释残差概念，不是新增实验。训练时先冻结源预测，再从目标记录学习其差异。如果源模型保留了主要规律，可以用源预测加修正；源结构不适合时，下一页所示的直接目标模型也可被选中。公交实际模型同时预测事件、状态、时间与乘客成本，不仅预测一个ETA。')
    deck.node(.9, 2.06, 3.03, 1.55, '源仿真预测', '下一事件间隔\n300秒', TEAL, 22)
    deck.node(5.12, 2.06, 3.03, 1.55, '目标历史记录', '对应实际间隔\n360秒', PURPLE, 22)
    deck.node(9.36, 2.06, 3.03, 1.55, '待学习的差异', '遗漏的延误\n+60秒', RED, 22)
    deck.line(3.93, 2.84, 5.0, 2.84, INK, 1.8, arrow=True)
    deck.line(8.15, 2.84, 9.24, 2.84, INK, 1.8, arrow=True)
    deck.text(.98, 4.13, 11.42, .65, '修正后的预测 = 保留下来的源规律 + 从目标记录学到的差异', 25, True, GOLD, PP_ALIGN.CENTER)
    deck.text(.98, 5.27, 11.42, .42, '示例秒数仅解释原理；真实模型同时预测事件、状态、时间与成本。', 17, color=MUTED, align=PP_ALIGN.CENTER)
    deck.verdict('有用的源结构可以复用；源结构不适合时，不强迫残差模型承担全部纠错。')


def confidence_visual():
    use(40, '模型预测已经有用，为什么训练时仍可能被忽略？', notes='本页不先讲递推公式，而用相同局部置信的解析示例说明问题。每一步模型预测的置信都为0.5，旧累计规则在传播折扣0.9下会给出约0.00098的权重，使模型样本几乎不参与价值学习。归一化后恒定置信对应的权重为0.5，恢复可用的模型更新强度。左边曲线与旧PPT相同，都是固定点的解析计算，不是实测学习曲线。下一页再看实际末期模型权重和公交服务结果，从而区分原理和证据。')
    deck.text(.83, 1.82, 11.6, .43, '每一步都给0.5置信：旧规则却把整条信息压到接近零', 23, True, RED)
    cs = [.1, .2, .3, .4, .5, .6, .7, .8, .9, 1]
    deck.chart(.85, 2.57, 7.48, 3.11, [str(x) for x in cs],
               [('原累计：c¹⁰', [x**10 for x in cs]), ('归一化：c', cs)],
               kind='line', colors=[RED, TEAL], ymin=0, ymax=1, number_format='0.0')
    deck.text(.91, 2.27, 7.13, .28, '横轴：局部置信；纵轴：实际模型样本权重（解析示例）', 14, color=MUTED)
    deck.text(8.95, 2.73, 3.45, .42, '相同局部置信 0.5', 23, True, INK)
    deck.text(8.95, 3.4, 3.45, .49, '原权重 0.00098', 25, True, RED)
    deck.text(8.98, 4.03, 3.45, .43, '模型几乎被忽略', 20, color=RED)
    deck.text(8.95, 4.76, 3.45, .49, '归一化权重 0.5', 25, True, TEAL)
    deck.text(8.98, 5.37, 3.45, .4, '可靠预测真正参与学习', 18, color=TEAL)
    deck.verdict('世界模型不仅要预测得较好，还需要在策略训练中保持有效作用。')


def write_outputs():
    # CS-BAPR pages remain contiguous and otherwise unchanged by request.
    backups = [9, 12, 18, 19, 29, 38]
    order = []
    for key in range(1, 49):
        if key not in backups:
            order.append(key)
        if key == 7:
            order.append('variance')
    main_count = len(order)
    order.extend(range(49, 55))
    order.extend(backups)
    ids = {key: item for key, item in zip(range(1, 55), list(deck.prs.slides._sldIdLst)[:54])}
    ids['variance'] = list(deck.prs.slides._sldIdLst)[-1]
    for item in list(deck.prs.slides._sldIdLst):
        deck.prs.slides._sldIdLst.remove(item)
    for key in order:
        deck.prs.slides._sldIdLst.append(ids[key])
    main_records = [records[key] for key in order[:main_count]]
    total = sum(r['seconds'] or 50 for r in main_records)
    times = [round((r['seconds'] or 50)*2400/total) for r in main_records]
    times[-2] += 2400-sum(times)
    elapsed = 0
    final = []
    for index, key in enumerate(order, 1):
        record = records[key]
        record['original_page'] = key
        record['page'] = index
        record['backup'] = index > main_count
        slide = slides[key]
        slide.shapes[5].text_frame.paragraphs[0].runs[0].text = f'{index:02d}'
        if record['backup']:
            if not record['title'].startswith('备用'):
                record['title'] = '备用｜' + record['title']
                slide.shapes[2].text_frame.paragraphs[0].runs[0].text = record['title']
                slide.shapes[2].text_frame.paragraphs[0].runs[0].font.size = Pt(26)
            prefix = '备用页，不计入主讲。'
            record['seconds'] = 0
        else:
            seconds = times[index-1]
            record['seconds'] = seconds
            record['start_seconds'] = elapsed
            elapsed += seconds
            record['end_seconds'] = elapsed
            prefix = f'建议讲述时间：{seconds}秒。'
        slide.notes_slide.notes_text_frame.text = prefix + '\n' + record['notes'] + '\n来源：' + record['source']
        for shape in slide.shapes:
            for ref in shape._element.xpath('.//a:effectRef'):
                ref.set('idx', '0')
        final.append(record)
    deck.prs.save(OUTPUT)
    (HERE / 'slides.json').write_text(json.dumps(final, ensure_ascii=False, indent=2), encoding='utf-8')
    lines = ['# 博士毕业答辩讲稿', '', f'主讲{main_count}页，约40分钟；其余{len(final)-main_count}页为备用。', '',
             '依据：《排版完成_残留问题替换_20261011.pdf》、现行答辩材料及各方法原论文。', '']
    for r in final:
        if r['backup']:
            stamp = '备用'
        else:
            stamp = f"{r['start_seconds']//60:02d}:{r['start_seconds']%60:02d}–{r['end_seconds']//60:02d}:{r['end_seconds']%60:02d}"
        lines.extend([f"## {r['page']:02d} {r['title']}（{stamp}）", '', r['notes'], '', '来源：'+r['source'], ''])
    (ROOT / 'output/博士毕业答辩_问题图解_讲稿_20261011.md').write_text('\n'.join(lines), encoding='utf-8')
    print(json.dumps({'pptx':str(OUTPUT),'slides':len(final),'main_slides':main_count,'main_seconds':elapsed,
                      'page_mapping':{str(r['original_page']):r['page'] for r in final}}, ensure_ascii=False))


if __name__ == '__main__':
    if len(sys.argv)>1 and sys.argv[1]=='assets':
        extract_trajectory()
        raise SystemExit
    HERE.mkdir(exist_ok=True)
    backup = HERE / 'before_edit'
    backup.mkdir(exist_ok=True)
    for name in [SOURCE.name, '博士毕业答辩_批注修订_预览_20261007.pdf', '博士毕业答辩_批注修订_讲稿_20261007.md']:
        target = backup / name
        if not target.exists():
            copy2(ROOT / 'output' / name, target)
    deck = Deck()
    deck.prs = Presentation(SOURCE)
    slides = {i+1:s for i,s in enumerate(deck.prs.slides)}
    records = {i+1:deepcopy(r) for i,r in enumerate(json.loads((OLD/'slides.json').read_text()))}
    reward_slide()
    resac_problem()
    variance_slide()
    resac_mechanism()
    bapr_problem()
    bapr_identification()
    residual_visual()
    confidence_visual()
    for key, title in [(28, 'H2O+：有目标记录，却不能在目标环境反复试错'),
                       (34, 'MC-WM：源仿真标定不准，策略会学到偏差')]:
        records[key]['title'] = title
        slides[key].shapes[2].text_frame.paragraphs[0].runs[0].text = title
    # Avoid references to moved formula pages in the main speaker notes.
    records[10]['notes'] = records[10]['notes'].replace('下一页', '后续')
    records[17]['notes'] = records[17]['notes'].replace('具体公式及一组数值示例放在下一页', '具体公式及一组数值示例放在备用页')
    records[20]['notes'] = records[20]['notes'].replace('零正则RE-SAC', '零正则RE-SAC（RE-SAC-0）')
    records[28]['notes'] = records[28]['notes'].replace('下一页', '后续')
    write_outputs()
