"""Build the Chinese thesis working draft with native Word equations."""

from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from xml.etree import ElementTree
import json
import re

import pypandoc
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from thesis_crossrefs import add_numbering, prepare_text


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "manuscript"
OUTPUT = ROOT / "output"
TITLE = "运行不确定性与环境失配下的公交在途学习控制方法研究"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def font(style, size=12, east_asia="宋体", latin="Times New Roman"):
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)


def set_layout(doc):
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        # Dimensions follow the supplied same-college thesis example.
        section.top_margin = Pt(72)
        section.bottom_margin = Pt(72)
        section.left_margin = Pt(89.85)
        section.right_margin = Pt(89.85)
        section.header_distance = Pt(42.55)
        section.footer_distance = Pt(49.6)
        section.different_first_page_header_footer = True


def set_styles(doc):
    if "Verbatim Char" not in doc.styles:
        doc.styles.add_style("Verbatim Char", WD_STYLE_TYPE.CHARACTER)
    for name in ("Normal", "Body Text", "First Paragraph"):
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        font(style)
        pf = style.paragraph_format
        pf.first_line_indent = Pt(24)
        pf.line_spacing = Pt(20)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.widow_control = True
    for level, size, before, after in ((1, 16, 12, 18), (2, 14, 10, 8), (3, 12, 8, 6)):
        style = doc.styles[f"Heading {level}"]
        font(style, size, "黑体")
        style.font.bold = True
        pf = style.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = Pt(20)
        pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        pf.keep_with_next = True
        pf.page_break_before = level == 1
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    for name, size in (("Title", 22), ("Subtitle", 12)):
        font(doc.styles[name], size, "黑体" if name == "Title" else "宋体")
        doc.styles[name].paragraph_format.first_line_indent = Pt(0)
        doc.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ("Verbatim Char", "Source Code"):
        if name in doc.styles:
            font(doc.styles[name], 9, latin="Consolas")
    if "Thesis Figure Caption" not in doc.styles:
        doc.styles.add_style("Thesis Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    style = doc.styles["Thesis Figure Caption"]
    font(style, 10.5)
    pf = style.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.line_spacing = Pt(16)
    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    pf.space_before = Pt(4)
    pf.space_after = Pt(6)
    pf.keep_together = True
    if "Bibliography" not in doc.styles:
        doc.styles.add_style("Bibliography", WD_STYLE_TYPE.PARAGRAPH)
    style = doc.styles["Bibliography"]
    font(style, 10.5)
    pf = style.paragraph_format
    pf.left_indent = Pt(24)
    pf.first_line_indent = Pt(-24)
    pf.line_spacing = Pt(17)
    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    pf.space_after = Pt(4)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.keep_together = True
    if "TOC Heading" not in doc.styles:
        doc.styles.add_style("TOC Heading", WD_STYLE_TYPE.PARAGRAPH)
    style = doc.styles["TOC Heading"]
    font(style, 16, "黑体")
    style.font.bold = True
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.page_break_before = True
    style.paragraph_format.space_after = Pt(18)
    style.paragraph_format.keep_with_next = True
    for level in range(1, 4):
        name = f"TOC {level}"
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        font(doc.styles[name])
        pf = doc.styles[name].paragraph_format
        pf.first_line_indent = Pt(0)
        pf.left_indent = Pt(12 * (level - 1))
        pf.line_spacing = Pt(20)
        pf.space_after = Pt(0)


def add_field(paragraph, instruction):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    paragraph._p.append(field)


def add_front_matter_structure(doc):
    abstract = next(p for p in doc.paragraphs if p.text == "摘要")
    chapter_one = next(p for p in doc.paragraphs if p.text.startswith("第1章"))
    notation = next((p for p in doc.paragraphs if p.text == "主要符号说明"), chapter_one)
    notation.insert_paragraph_before("目录", "TOC Heading")
    toc = notation.insert_paragraph_before()
    toc.paragraph_format.first_line_indent = Pt(0)
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    # Separate cover, preliminary pages, and the Arabic-numbered thesis body.
    for boundary in (abstract, chapter_one):
        paragraph = boundary.insert_paragraph_before()
        paragraph._p.get_or_add_pPr().append(deepcopy(doc._element.body.sectPr))
        boundary.paragraph_format.page_break_before = False
    for index, section in enumerate(doc.sections):
        section.different_first_page_header_footer = index == 0
        number = OxmlElement("w:pgNumType")
        number.set(qn("w:fmt"), "upperRoman" if index == 1 else "decimal")
        number.set(qn("w:start"), "1")
        section._sectPr.append(number)


def format_document(doc):
    set_layout(doc)
    set_styles(doc)
    add_front_matter_structure(doc)
    for section in doc.sections:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        header = section.header.paragraphs[0]
        header.text = TITLE
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header.paragraph_format.first_line_indent = Pt(0)
        for run in header.runs:
            run.font.size = Pt(9)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.first_line_indent = Pt(0)
        add_field(footer, "PAGE")
    # Exact body leading can clip matrix and fraction structures in OMML.
    paragraphs = doc.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.style.name == "Bibliography" and re.search(r"[\u4e00-\u9fff]", paragraph.text):
            for node in paragraph._p.xpath(".//w:t"):
                if node.text:
                    node.text = node.text.replace("et al.", "等.")
        if paragraph.style.name == "Title":
            paragraph.text = TITLE.replace("的公交", "的\n公交", 1)
        if paragraph._p.xpath(".//w:drawing"):
            pf = paragraph.paragraph_format
            pf.first_line_indent = Pt(0)
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.space_before = Pt(8)
            pf.space_after = Pt(0)
            pf.keep_with_next = True
        if paragraph.text.startswith(tuple(f"图{chapter}-" for chapter in [*range(1, 8), "F"])):
            prefix = paragraph.text.split(" ", 1)[0]
            if prefix.rsplit("-", 1)[-1].isdigit():
                paragraph.style = doc.styles["Thesis Figure Caption"]
                if index + 1 < len(paragraphs) and not paragraphs[index + 1].style.name.startswith("Heading"):
                    paragraph.paragraph_format.keep_with_next = True
                    paragraphs[index + 1].paragraph_format.keep_together = True
        if any(run.style.name == "Verbatim Char" for run in paragraph.runs):
            paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if paragraph._p.xpath(".//m:oMath"):
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            paragraph.paragraph_format.line_spacing = Pt(20)
        if paragraph._p.xpath(".//m:oMathPara"):
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(5)
        if paragraph.text.startswith(("定理", "引理", "命题", "推论", "反例")):
            if index + 1 < len(paragraphs) and paragraphs[index + 1]._p.xpath(".//m:oMathPara"):
                paragraph.paragraph_format.keep_with_next = True
        if paragraph.text.startswith(tuple(f"表{chapter}-" for chapter in [*range(1, 8), "F"])):
            prefix = paragraph.text.split(" ", 1)[0]
            if prefix.rsplit("-", 1)[-1].isdigit():
                paragraph.paragraph_format.keep_with_next = True
    width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for table in doc.tables:
        table.autofit = False
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single" if edge in ("top", "bottom") else "nil")
            border.set(qn("w:sz"), "8")
            border.set(qn("w:color"), "000000")
            borders.append(border)
        table._tbl.tblPr.append(borders)
        column_count = len(table.columns)
        fractions = {
            2: (0.32, 0.68),
            3: (0.23, 0.43, 0.34),
            4: (0.40, 0.20, 0.20, 0.20),
            5: (0.32, 0.17, 0.17, 0.17, 0.17),
            6: (0.14, 0.26, 0.15, 0.15, 0.15, 0.15),
            7: (0.28, 0.12, 0.12, 0.12, 0.12, 0.12, 0.12),
        }[column_count]
        if column_count == 3 and table.cell(0, 0).text in ("方法", "控制器"):
            fractions = (0.30, 0.35, 0.35)
        if column_count == 3 and table.cell(0, 0).text == "指标":
            fractions = (0.46, 0.27, 0.27)
        if column_count == 5 and table.cell(0, 1).text == "历史训练":
            fractions = (0.28, 0.15, 0.21, 0.15, 0.21)
        if column_count == 4 and table.cell(0, 1).text in ("数据及交互条件", "决策过程与动作"):
            fractions = (0.18, 0.27, 0.25, 0.30)
        if column_count == 5 and table.cell(0, 0).text == "项目":
            fractions = (0.18, 0.205, 0.205, 0.205, 0.205)
        if column_count == 5 and table.cell(0, 0).text == "状态分组":
            fractions = (0.25, 0.12, 0.27, 0.18, 0.18)
        if column_count == 6 and table.cell(0, 0).text == "基础策略":
            fractions = (0.20, 0.14, 0.14, 0.16, 0.18, 0.18)
        if column_count == 5 and table.cell(0, 1).text == "评论家数量":
            fractions = (0.22, 0.12, 0.22, 0.22, 0.22)
        if column_count == 6 and table.cell(0, 1).text == "需求倍率":
            fractions = (0.19, 0.10, 0.17, 0.18, 0.17, 0.19)
        if column_count == 7 and table.cell(0, 0).text == "训练重复":
            fractions = (0.10, 0.115, 0.115, 0.12, 0.18, 0.18, 0.19)
        for column, fraction in zip(table.columns, fractions):
            column.width = int(width * fraction)
        for row_index, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for cell, fraction in zip(row.cells, fractions):
                cell.width = int(width * fraction)
                if row_index == 0:
                    cell_borders = OxmlElement("w:tcBorders")
                    bottom = OxmlElement("w:bottom")
                    bottom.set(qn("w:val"), "single")
                    bottom.set(qn("w:sz"), "4")
                    cell_borders.append(bottom)
                    cell._tc.get_or_add_tcPr().append(cell_borders)
                for paragraph in cell.paragraphs:
                    pf = paragraph.paragraph_format
                    pf.first_line_indent = Pt(0)
                    pf.line_spacing = Pt(15)
                    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
                    pf.space_before = Pt(3)
                    pf.space_after = Pt(3)
                    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    # The thesis tables are short; keep their rows on one page.
                    pf.keep_with_next = row_index < len(table.rows) - 1 and table.cell(0, 0).text != "符号"
                    for run in paragraph.runs:
                        run.font.size = Pt(9 if run.style.name == "Verbatim Char" else 10.5)
                        run.font.bold = row_index == 0
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = TITLE
    doc.core_properties.author = ""


def main():
    OUTPUT.mkdir(exist_ok=True)
    front = sorted((SOURCE / "frontmatter").glob("*.md"))
    chapters = sorted((SOURCE / "chapters").glob("*.md"))
    appendices = sorted(path for path in (SOURCE / "appendices").glob("*.md")
                        if path.name[0] in "ABCDE" or path.name == "F_公交世界模型的探索性比较.md")
    sources = front + chapters + appendices
    content = "\n\n".join(
        [path.read_text(encoding="utf-8") for path in front + chapters]
        + ["# 参考文献\n\n::: {#refs}\n:::"]
        + [path.read_text(encoding="utf-8") for path in appendices]
    )
    content, equation_records, caption_targets = prepare_text(content)
    ast = json.loads(pypandoc.convert_text(content, "json", format="markdown+tex_math_dollars"))
    cited = set()

    def collect_citations(node):
        if isinstance(node, dict):
            if node.get("t") == "Cite":
                cited.update(c["citationId"] for c in node["c"][0])
            for value in node.values():
                collect_citations(value)
        elif isinstance(node, list):
            for value in node:
                collect_citations(value)

    collect_citations(ast)
    references = json.loads((SOURCE / "references.json").read_text(encoding="utf-8"))
    missing = cited - {entry["id"] for entry in references}
    assert not missing, f"Missing bibliography entries: {sorted(missing)}"
    merged = OUTPUT / "博士论文_正文与数学附录_阶段稿.md"
    merged.write_text(content, encoding="utf-8")
    reference = OUTPUT / "word_reference.docx"
    seed = Document()
    set_layout(seed)
    set_styles(seed)
    seed.save(reference)
    result = OUTPUT / "博士论文_正文与数学附录_阶段稿.docx"
    metadata = (
        "---\n"
        f'title: "{TITLE}"\n'
        'subtitle: "博士学位论文"\n'
        'lang: zh-CN\n'
        'toc-title: "目录"\n'
        "---\n\n"
    )
    pypandoc.convert_text(
        metadata + content,
        "docx",
        format="markdown+tex_math_dollars",
        outputfile=str(result),
        extra_args=[
            "--standalone", "--citeproc", "--metadata=link-citations:true", f"--reference-doc={reference}",
            f"--resource-path={SOURCE}",
            f"--bibliography={SOURCE / 'references.json'}",
            f"--csl={SOURCE / 'gb-t-7714-2015-numeric.csl'}",
        ],
    )
    document = Document(result)
    format_document(document)
    numbered_equations = add_numbering(document, equation_records, caption_targets)
    document.save(result)
    # A missing OMML count means formula conversion failed, not a successful build.
    with ZipFile(result) as archive:
        tree = ElementTree.fromstring(archive.read("word/document.xml"))
        formula_count = len(tree.findall(f".//{{{M}}}oMath"))
        display_count = len(tree.findall(f".//{{{M}}}oMathPara"))
        media_count = sum(name.startswith("word/media/") for name in archive.namelist())
    expected_math = len(re.findall(r"\$\$.*?\$\$|(?<!\$)\$(?!\$).*?(?<!\$)\$(?!\$)", content, re.S))
    assert formula_count == expected_math, (formula_count, expected_math)
    assert display_count == content.count("\n$$\n") // 2
    assert media_count == content.count("![]("), "A thesis figure was not embedded"
    report = {
        "document": str(result),
        "sources": [str(path.relative_to(ROOT)) for path in sources],
        "native_equations": formula_count,
        "display_equations": display_count,
        "numbered_equations": numbered_equations,
        "embedded_media": media_count,
        "source_characters": len(content),
        "cited_references": len(cited),
        "visual_layout_verified": False,
    }
    (OUTPUT / "word_build.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (OUTPUT / "equation_index.json").write_text(json.dumps(equation_records, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
