"""Check the requested preservation boundaries and native-object content."""
from copy import deepcopy
from pathlib import Path
import json
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SOURCE = ROOT / 'output/博士毕业答辩_批注修订_20261007.pptx'
OUT = ROOT / 'output/博士毕业答辩_问题图解_20261011.pptx'
before, after = Presentation(SOURCE), Presentation(OUT)
records = json.loads((HERE / 'slides.json').read_text())
mapping = {r['original_page']:r['page'] for r in records}
issues = []
for i, (slide, record) in enumerate(zip(after.slides, records), 1):
    if record['title'] not in '\n'.join(s.text for s in slide.shapes if s.has_text_frame):
        issues.append(f'{i}: missing title')
    if record['notes'] not in slide.notes_slide.notes_text_frame.text:
        issues.append(f'{i}: missing notes')
    for s in slide.shapes:
        if min(s.left,s.top)<0 or s.left+s.width>after.slide_width+20 or s.top+s.height>after.slide_height+20:
            issues.append(f'{i}: off-slide shape {s.name}')

def shapes_except_page(slide):
    xml=deepcopy(slide._element)
    for shape in xml.xpath('./p:cSld/p:spTree/p:sp'):
        if shape.xpath('./p:nvSpPr/p:cNvPr[@id="7"]'):
            shape.getparent().remove(shape)
    return etree.tostring(xml)

cs_preservation={}
for i in range(21,28):
    cs_preservation[str(i)]=shapes_except_page(before.slides[i-1])==shapes_except_page(after.slides[mapping[i]-1])
    if not cs_preservation[str(i)]:
        issues.append(f'Original CS slide {i}: slide geometry/content changed')
old_chart=next(s.chart for s in before.slides[21].shapes if s.has_chart)
new_chart=next(s.chart for s in after.slides[mapping[22]-1].shapes if s.has_chart)

def response_data(chart):
    return [([v.text for v in s._element.xpath('./c:xVal/c:numRef/c:numCache/c:pt/c:v')], list(s.values)) for s in list(chart.series)[1:]]
response_unchanged=response_data(old_chart)==response_data(new_chart)
cs_chart_unchanged=etree.tostring(old_chart._element)==etree.tostring(new_chart._element)
sample_count=len(new_chart.series[0].values)
filter_old=next(s for s in before.slides[16].shapes if s.shape_type==13)
filter_new=next(s for s in after.slides[mapping[17]-1].shapes if s.shape_type==13)
with ZipFile(OUT) as z:
    math_count=sum(len(etree.fromstring(z.read(f'ppt/slides/slide{i}.xml')).xpath('//*[local-name()="oMath"]')) for i in range(1,56))
result={
    'slides':len(after.slides),'main_slides':sum(not r['backup'] for r in records),
    'main_seconds':sum(r['seconds'] for r in records), 'native_equations':math_count,
    'main_native_equations':sum(len(s._element.xpath('//*[local-name()="oMath"]')) for s in list(after.slides)[:43]),
    'original_pptx_unchanged':SOURCE.read_bytes()==(HERE/'before_edit'/SOURCE.name).read_bytes(),
    'CS_BAPR_slide_content_and_geometry_preserved_except_page_number':cs_preservation,
    'CS_response_curves_unchanged':response_unchanged,
    'CS_complete_chart_unchanged':cs_chart_unchanged,
    'CS_sample_points':sample_count,'CS_sample_marker_size':new_chart.series[0].marker.size,
    'original_filter_image_preserved':filter_old.image.blob==filter_new.image.blob,
    'issues':issues,
}
(HERE/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False,indent=2))
assert not issues and response_unchanged and cs_chart_unchanged and result['original_pptx_unchanged'] and result['original_filter_image_preserved']
assert len(after.slides)==55 and result['main_seconds']==2400 and sample_count==len(old_chart.series[0].values)
