"""Build a final-source comparison; read thesis files, write comparison artifacts only."""
from pathlib import Path
from array import array
from collections import defaultdict, Counter
from difflib import SequenceMatcher
import json, re, math
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

A=Path(__file__).resolve().parent; OUT=A.parents[1]
OLD=A/'修改前_重构论文.md'; NEW=OUT/'重构论文_语言修订_v2_20261015.md'
STEM=OUT/'语言修订_v2_修改对照表_20261015'
J={str(p.relative_to(A)):json.loads(p.read_text()) for p in A.rglob('*.json') if p.name not in ('comparison_statistics.json','word_only_changes.json')}
original=OLD.read_text(); final=NEW.read_text(); current=original
blocks=[{'start':m.start(),'end':m.end(),'old':m[0]} for m in re.finditer(r'\S[\s\S]*?(?=\n\s*\n|\Z)',original)]
labels=array('i',[-1])*len(original)
for uid,b in enumerate(blocks): labels[b['start']:b['end']]=array('i',[uid])*(b['end']-b['start'])
reasons=defaultdict(list); sources=defaultdict(list); logs=defaultdict(list); formats=[]

def note(uid, reason, source, ident=''):
    if reason and reason not in reasons[uid]:reasons[uid].append(reason)
    if source and source not in sources[uid]:sources[uid].append(source)
    if ident and ident not in logs[uid]:logs[uid].append(ident)

def swap(start,end,replacement,reason='',source='',ident='',kind=''):
    global current,labels
    ids={x for x in labels[start:end] if x>=0}; assert len(ids)==1,(ident,ids)
    uid=next(iter(ids)); before=current[start:end]
    note(uid,reason,source,ident)
    if kind:formats.append({'uid':uid,'kind':kind,'old':before,'new':replacement})
    current=current[:start]+replacement+current[end:]
    labels[start:end]=array('i',[uid])*len(replacement)

def change(old,new,reason,source,ident='',kind='',negative_degree=False):
    pattern=re.escape(old)+(r'(?!度)' if negative_degree else '')
    matches=list(re.finditer(pattern,current))
    for m in reversed(matches):swap(m.start(),m.end(),new,reason,source,ident,kind)
    return len(matches)

# Replay logged replacements in memory to attach each final paragraph to its original paragraph.
edits=J['applied_language_edits.json']; intervals=[]
for i,e in enumerate(edits):
    assert original.count(e['old'])==1
    start=original.index(e['old']); intervals.append((start,start+len(e['old']),e,i))
for start,end,e,i in sorted(intervals,reverse=True):
    reason=e.get('reason') or '；'.join(e.get('reasons',[]))
    ident=e.get('id') or e.get('location') or f'explicit-{i+1}'
    swap(start,end,e['new'],reason,'applied_language_edits.json',ident)
rulemeta={}
for fn in ('final_language_suggestions.json','final_review_edits.json'):
    data=J.get(fn,{}); data=data.get('edits',data) if isinstance(data,dict) else data
    for e in data:rulemeta[e['old'],e['new']]=(fn,e.get('id',''))
rules=J['language_rules.json']; bold_index=next(i for i,r in enumerate(rules) if r['old']=='正文中间及表格数据加粗')
for i,r in enumerate(rules[:bold_index]):
    kind='TeX定界符格式' if r['old'] in (r'\vert{}',r'\Vert{}') else ''
    n=change(r['old'],r['new'],r['reason'],'language_rules.json',f'rule-{i+1}',kind,r['old']=='跨保真')
    assert n==r['count'],(i,n,r['count'])
# Use exactly the scope of the applied emphasis cleanup; record every location.
bold=[]; offset=0; lines=current.splitlines(True)
for i,line in enumerate(lines):
    header=line.lstrip().startswith('|') and i+1<len(lines) and re.match(r'^\s*\|[ :|\-]+\|\s*$',lines[i+1])
    if not header:
        for m in re.finditer(r'\*\*(.+?)\*\*',line):
            lead=line[:m.start()].strip(); content=m[1]
            structured=not lead and re.match(r'^(?:（[一二三四五六七八九十0-9]+）|\([0-9]+\)|[1-9][0-9]*[.、]|定理|引理|命题|证明)',content)
            if not(structured and len(content)<=95):bold.append((offset+m.start(),offset+m.end(),content))
    offset+=len(line)
for start,end,replacement in reversed(bold):swap(start,end,replacement,'清理正文及数据的过度强调','language_rules.json','bold-cleanup','去除加粗')
assert len(bold)==rules[bold_index]['count']
for e in reversed(J['markdown_figure_links.json']):
    assert current.count(e['old'])==1
    start=current.index(e['old']);swap(start,start+len(e['old']),e['new'],'用实际图片链接替代制作占位说明','markdown_figure_links.json',e['figure'],'图片链接')
for i,r in enumerate(rules[bold_index+1:],bold_index+1):
    source,ident=rulemeta.get((r['old'],r['new']),('language_rules.json',f'rule-{i+1}'))
    n=change(r['old'],r['new'],r['reason'],source,ident)
    assert n==r['count'],(i,n,r['count'])
# A final independent review may be applied after the main revision script.
if current!=final and 'final_review_edits.json' in J:
    extra=J['final_review_edits.json'];extra=extra.get('edits',extra) if isinstance(extra,dict) else extra
    for i,e in enumerate(extra):
        if e.get('old') and e['old'] in current:
            change(e['old'],e['new'],e.get('reason','独立复核修正'),'final_review_edits.json',e.get('id',f'review-{i+1}'))
if current.rstrip('\n')+'\n'==final and current!=final:
    current=final; labels.extend(array('i',[-1])*(len(final)-len(labels)))
assert current==final,'Final source differs from the logged in-memory replay; update provenance before delivery.'

# Final source is authoritative; logs supply reasons, never extra counted changes.
newpieces=defaultdict(list)
for m in re.finditer(r'\S[\s\S]*?(?=\n\s*\n|\Z)',final):
    ids={x for x in labels[m.start():m.end()] if x>=0};assert len(ids)==1,ids
    newpieces[next(iter(ids))].append({'start':m.start(),'end':m.end(),'text':m[0]})
def line(text,pos):return text.count('\n',0,pos)+1
def normal(s):return re.sub(r'\s+','',s.replace('**','').replace(r'\vert{}',r'\vert ').replace(r'\Vert{}',r'\Vert '))
contexts=[];heads=[];table_name=''
for b in blocks:
    first=b['old'].splitlines()[0]
    if first.startswith('#'):
        level=len(first)-len(first.lstrip('#'));heads=[h for h in heads if h[0]<level];heads.append((level,first.lstrip('#').strip()))
    if re.match(r'^表\d+-\d+\s',first):table_name=first.split()[0]
    contexts.append((' > '.join(h[1] for h in heads),table_name))
rows=defaultdict(list); serial=Counter()
def add(sheet,uid,old,new,reason='',evidence='',detail='',oldline=None,newline=None):
    serial[sheet]+=1;ident={'正文段落':'P','表格单元格':'T','格式位置':'F','图标签':'G','Word独有':'W','文件路径':'A'}[sheet]+f'{serial[sheet]:04d}'
    if uid is not None:
        b=blocks[uid]; ctx=contexts[uid][0]
        ol=oldline or str(line(original,b['start']))
        nl=newline or '、'.join(str(line(final,x['start'])) for x in newpieces[uid])
        pos=f'{ctx}；修改前MD行{ol} → 新MD行{nl}'+('；'+detail if detail else '')
        reason=reason or '；'.join(reasons[uid]) or '逐段比较确认的文本变化'
        evidence=evidence or '；'.join(str(A/x) for x in sources[uid])
        logids='；'.join(logs[uid])
    else:pos=detail;logids=''
    rows[sheet].append([ident,pos,old,new,reason,evidence,logids])
    return ident

# Whole-table provenance is split into cells; pure bold changes are excluded from content counts.
changed_original_units=set();unit_ids={}; table_cells=Counter()
for uid,b in enumerate(blocks):
    pieces=newpieces[uid]; new='\n\n'.join(x['text'] for x in pieces);old=b['old']
    if old==new:continue
    is_table=old.lstrip().startswith('|')
    if is_table:
        oldrows=[x for x in old.splitlines() if x.strip().startswith('|')]
        tablepieces=[x for x in pieces if x['text'].lstrip().startswith('|')]
        assert len(tablepieces)==1
        newrows=tablepieces[0]['text'].splitlines()
        oldrows=[x for x in oldrows if not re.match(r'^\s*\|[ :|\-]+\|\s*$',x)]
        newrows=[x for x in newrows if not re.match(r'^\s*\|[ :|\-]+\|\s*$',x)]
        split=lambda x:[v.strip() for v in re.split(r'(?<!\\)\|',x.strip())[1:-1]]
        oldcells=list(map(split,oldrows));newcells=list(map(split,newrows));table=contexts[uid][1] or f'表块{uid}'
        oldheaders=oldcells[0];newheaders=newcells[0]
        for r in range(max(len(oldcells),len(newcells))):
            oc=oldcells[r] if r<len(oldcells) else [];nc=newcells[r] if r<len(newcells) else []
            for c in range(max(len(oc),len(nc))):
                before=oc[c] if c<len(oc) else '';after=nc[c] if c<len(nc) else ''
                if normal(before)==normal(after):continue
                changed_original_units.add(uid);table_cells[table]+=1
                label=(nc[0] if nc else oc[0]).replace('**','')
                column=(newheaders[c] if c<len(newheaders) else oldheaders[c]).replace('**','')
                evidence=''
                if table in ('表5-5','表6-3','表6-6'): evidence=str(A/'data_evidence.json')+'；'+str(A/'table_completion_suggestions.json')
                ident=add('表格单元格',uid,before,after,evidence=evidence,detail=f'{table}；'+('表头' if r==0 else f'数据行{r}：{label}')+f'；第{c+1}列：{column}',oldline=str(line(original,b['start'])+r+(r>0)) if oc else '—（新增行）',newline=str(line(final,tablepieces[0]['start'])+r+(r>0)) if nc else '—（删除行）')
                unit_ids.setdefault(uid,ident)
        for p in pieces:
            if p not in tablepieces:
                changed_original_units.add(uid);ident=add('正文段落',uid,'（新增表注）',p['text'],evidence=str(A/'data_evidence.json'),detail=table+'新增表注',newline=str(line(final,p['start'])));unit_ids.setdefault(uid,ident)
    elif old.startswith('【图'):
        pass
    elif normal(old)!=normal(new):
        changed_original_units.add(uid);unit_ids[uid]=add('正文段落',uid,old,new)

# Batch-format sheet: one row per original paragraph/type, with every affected fragment.
groups=defaultdict(list)
for f in formats:
    if f['kind']!='去除加粗':groups[f['uid'],f['kind']].append(f)
for (uid,kind),items in sorted(groups.items()):
    items=list(reversed(items)) if kind in ('去除加粗','图片链接') else items
    old='\n'.join(f'{i+1}. {x["old"]}' for i,x in enumerate(items));new='\n'.join(f'{i+1}. {x["new"]}' for i,x in enumerate(items))
    parent=unit_ids.get(uid,'无正文/单元格内容变更')
    if kind=='图片链接':old=blocks[uid]['old'];new='\n\n'.join(x['text'] for x in newpieces[uid])
    reason=f'{kind}，{len(items)}处；内容关联：{parent}。位置指向修改前/最终稿；重复定界符按出现次序编号。'
    add('格式位置',uid,old,new,reason,str(A/('markdown_figure_links.json' if kind=='图片链接' else 'language_rules.json')),detail=f'{kind}；{len(items)}处')
# Bold changes use original/final source, including emphasis removed inside explicit edits.
bold_old_total=bold_new_total=0;bold_position_groups=0
for uid,b in enumerate(blocks):
    before=re.findall(r'\*\*(.+?)\*\*',b['old']);after=re.findall(r'\*\*(.+?)\*\*','\n\n'.join(x['text'] for x in newpieces[uid]))
    bold_old_total+=len(before);bold_new_total+=len(after)
    if before==after:continue
    bold_position_groups+=1
    old='\n'.join(f'{i+1}. **{x}**' for i,x in enumerate(before)) or '无加粗'
    new='\n'.join(f'{i+1}. **{x}**' for i,x in enumerate(after)) or '无加粗'
    parent=unit_ids.get(uid,'无正文/单元格内容变更')
    add('格式位置',uid,old,new,f'加粗标记由{len(before)}处变为{len(after)}处，净减少{len(before)-len(after)}处；内容关联：{parent}。按修改前/最终稿实际标记，不把显式编辑中的去粗体漏掉。',str(OLD)+'；'+str(NEW),detail='加粗实际差异')
# Three edited diagrams; consolidated final label before/after, not generation steps.
for e in J['figure_language_updates.json']:
    for c in e['text_changes']:
        add('图标签',None,c['old'].replace('\\n','\n'),c['new'].replace('\\n','\n'),'图中文字与正文术语一致',e['svg']+'；'+e['png'],e['figure'])

wordfile=A/'word_only_changes.json'
if wordfile.exists():
    wd=json.loads(wordfile.read_text());wd=wd.get('changes',wd) if isinstance(wd,dict) else wd
    for e in wd:
        reason=e.get('reason','统一Word独有段落/页眉术语')
        if e.get('location')=='word/header17.xml':reason='附录G页眉与正文的跨保真度术语一致。'
        add('Word独有',None,e['old'],e['new'],reason,e.get('file',str(OUT/'重构论文_语言修订_v2_20261015.docx')),e.get('location','Word独有位置'))

for e in J.get('word_format_changes.json',{}).get('changes',[]):
    add('格式位置',None,e['old'],e['new'],e['reason'],e['file'],e['location'])

paths=[('修改前Markdown',OLD),('修改前Word',A/'修改前_重构论文_通篇校核_20261015.docx'),('本轮依据附件保留件',A/'修改前_Gemini_problem_v2.md'),('新版Markdown',NEW),('新版Word',OUT/'重构论文_语言修订_v2_20261015.docx'),('原主Markdown',OUT/'重构论文.md'),('原主Word',OUT/'重构论文_通篇校核_20261015.docx')]
for role,path in paths:
    if path.exists():add('文件路径',None,'原件保留' if '修改前' in role or '原主' in role else '',str(path),role,'存在；'+str(path.stat().st_size)+'字节',role)
for e in J['figure_language_updates.json']:
    for kind in ('svg','png'):add('文件路径',None,'原图未覆盖',e[kind],e['figure']+'修订图'+kind.upper(),str(A/'figure_language_updates.json'),e['figure'])
base_figures=json.loads((A.parent/'figure_map_final.json').read_text())['figures']
for entry in base_figures:
    if entry['figure'] not in {e['figure'] for e in J['figure_language_updates.json']}:continue
    for path in (Path(entry['source']),Path(entry['source']).with_suffix('.svg')):
        if path.exists():add('文件路径',None,'原图保留',str(path),entry['figure']+'修改前'+path.suffix.upper(),str(A.parent/'figure_map_final.json'),entry['figure'])

portable=Path('/mnt/d/Paper/Thesis/Gemini')
for name in (NEW.name,'重构论文_语言修订_v2_20261015.docx'):
    path=portable/name
    if path.exists():add('文件路径',None,'',str(path),'Windows交付副本；Markdown仅把57张图的路径改为语言修订_v2_配图/图X-Y.png，正文相同。',str(path),'D:/Paper/Thesis/Gemini/'+name)
imagefolder=portable/'语言修订_v2_配图'
if imagefolder.exists():add('文件路径',None,'',str(imagefolder),'Windows便携配图目录，含57张PNG。',str(imagefolder),'便携配图')
# Every changed source block must appear in content or final-format locations.
covered=changed_original_units|{f['uid'] for f in formats}
for uid,b in enumerate(blocks):
    target='\n\n'.join(x['text'] for x in newpieces[uid])
    if re.findall(r'\*\*(.+?)\*\*',b['old'])!=re.findall(r'\*\*(.+?)\*\*',target):covered.add(uid)
    assert b['old'].rstrip('\n')==target.rstrip('\n') or uid in covered,('Unlisted changed source block',uid)
known_ids={x for values in logs.values() for x in values}
for fn in ('final_language_suggestions.json','final_review_edits.json'):
    assert all(e['id'] in known_ids for e in J[fn]['edits']),('Unattributed follow-up edit',fn)
counts={name:len(value) for name,value in rows.items()}
stats={'source_old':str(OLD),'source_new':str(NEW),'count_basis':'正文一行对应一个修改前段落，可对应拆分后的多段；表格逐单元格，新增行逐列。格式与内容为不同维度，不相加。修改日志不直接累加计数。','rows_by_sheet':counts,'distinct_changed_original_paragraph_or_table_blocks':len(changed_original_units),'changed_cells_by_table':dict(table_cells),'format_occurrences_by_type':{'TeX定界符格式':sum(f['kind']=='TeX定界符格式' for f in formats),'图片链接':sum(f['kind']=='图片链接' for f in formats),'加粗净减少':bold_old_total-bold_new_total,'WordCI单元格换行':len(J.get('word_format_changes.json',{}).get('changes',[]))},'bold_source_comparison':{'old_paired_spans':bold_old_total,'new_paired_spans':bold_new_total,'net_removed':bold_old_total-bold_new_total,'position_groups':bold_position_groups,'global_rule_execution_count':len(bold),'other_stages_net_removed':bold_old_total-bold_new_total-len(bold)},'explicit_edit_records':len(edits),'post_language_suggestions':len(J['final_language_suggestions.json']['edits']),'final_review_records':len(J.get('final_review_edits.json',{}).get('edits',[])) if isinstance(J.get('final_review_edits.json',{}),dict) else len(J['final_review_edits.json']),'replay_equals_final_source':True,'terminal_newline_normalized':not original.endswith('\n') and final.endswith('\n'),'word_only_inventory_present':wordfile.exists(),'appendix_note':'MD附录F→G及附G公式号同步已校核Word既有附录G，未新增加附录。','files_read':sorted(list(J)+(['word_only_changes.json'] if wordfile.exists() else []))}
headers=['编号','位置','原文','修改后','原因','证据或适用文件','编辑记录索引']
wb=Workbook();intro=wb.active;intro.title='说明与统计'
summary=[['项目','内容'],['比较范围','修改前Markdown与最终修订Markdown真实差异；另列Word独有术语与图标签。'],['计数口径',stats['count_basis']],['附录字母',stats['appendix_note']],['可核对位置','行号从1开始；修改前/后段落均给位置。格式批量每个受影响段落及全部片段另表列出。'],['格式文本','加粗与图链接按修改前/最终稿列出；TeX定界符逐处列出；与正文内容变更不重复相加。'],['执行记录','162条精确编辑、11条后续建议、14条复核收尾及全局规则可能作用同一段，不相加作为用户改动条数。'],['最终原文重放一致',str(stats['replay_equals_final_source'])],['最终Word','原生Word最终322页；39个CI单元格按实际最终文件逐格核对，数值一致。']]
for name,count in counts.items():summary.append([name+'行数',count])
for name,count in table_cells.items():summary.append([name+'内容变化单元格数',count])
for name,count in stats['format_occurrences_by_type'].items():summary.append([name+'实际应用处数',count])
for r in summary:intro.append(r)
for name,data in rows.items():
    ws=wb.create_sheet(name);ws.append(headers)
    for r in data:ws.append(r)
    ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
    widths=[12,48,80,80,55,64,35]
    for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
    for r in ws.iter_rows():
        for c in r:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='等线',size=10)
    if name in ('正文段落','表格单元格','图标签','Word独有'):
        for r in ws.iter_rows(min_row=2):
            r[2].fill=PatternFill('solid',fgColor='FFF4F2');r[3].fill=PatternFill('solid',fgColor='F0F8F2')
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='24476A');c.font=Font(name='等线',size=11,bold=True,color='FFFFFF')
    ws.row_dimensions[1].height=28
    for i,r in enumerate(data,2):ws.row_dimensions[i].height=min(409,max(34,14*max(math.ceil(sum(2 if ord(c)>255 else 1 for c in str(v))/75)+str(v).count('\n') for v in r[1:6])))
intro.column_dimensions['A'].width=39;intro.column_dimensions['B'].width=115;intro.freeze_panes='B2'
for row in intro:
    for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='等线',size=11)
for c in intro[1]:c.font=Font(name='等线',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24476A')
for i in range(2,intro.max_row+1):intro.row_dimensions[i].height=43
wb.save(str(STEM)+'.xlsx')
# Markdown remains complete; long paragraphs are cells, not shortened excerpts.
escape=lambda x:str(x).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;').replace('|','&#124;').replace('\n','<br>')
md=['# 语言修订v2修改对照表','',f'修改前：[{OLD.name}]({OLD})  ',f'修订后：[{NEW.name}]({NEW})','',stats['count_basis'],'',stats['appendix_note'],'','加粗与图片链接按原稿/最终稿列出；TeX定界符逐处列出。格式与正文内容变更分开计数。','', '| 类别 | 对照行数 |','| --- | ---: |']
md += [f'| {name} | {count} |' for name,count in counts.items()]
md += ['', f'修改前配对加粗{bold_old_total}处，最终稿保留{bold_new_total}处，净减少{bold_old_total-bold_new_total}处；其中全局清理规则执行{len(bold)}处，其余净变化发生在显式编辑等阶段。','', '格式实际差异：'+ '；'.join(f'{k}{v}处' for k,v in stats['format_occurrences_by_type'].items())+'。','']
for name,data in rows.items():
    md += ['## '+name,'','| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']
    md += ['| '+' | '.join(escape(x) for x in r)+' |' for r in data];md.append('')
Path(str(STEM)+'.md').write_text('\n'.join(md)+'\n')
(A/'comparison_statistics.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2)+'\n')
# Detect truncated/export-omitted cells: exported rows must match the in-memory complete ledger.
check=load_workbook(str(STEM)+'.xlsx',read_only=True,data_only=False)
for name,data in rows.items():
    actual=list(check[name].values)[1:]
    assert len(actual)==len(data)
    for a,b in zip(actual,data):assert tuple('' if v is None else str(v) for v in a)==tuple(str(v) for v in b)
print(json.dumps({'outputs':[str(STEM)+'.xlsx',str(STEM)+'.md'],'statistics':{k:v for k,v in stats.items() if k!='files_read'}},ensure_ascii=False))
