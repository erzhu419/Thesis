from pathlib import Path
from shutil import copy2

import fitz
from pptx import Presentation

here = Path(__file__).resolve().parent
out = here.parents[1] / 'output'
assets = here / 'comment_assets'
backup = here / 'before_comments'
assets.mkdir(exist_ok=True)
backup.mkdir(exist_ok=True)
for path in [
    out / '博士毕业答辩_20261007.pptx',
    out / '博士毕业答辩_预览_20261007.pdf',
    out / '博士毕业答辩_讲稿_20261007.md',
    here / 'build_defense.py', here / 'drawing.py', here / 'check_deck.py',
    here / 'slides.json', here / 'verification.json',
]:
    target = backup / path.name
    if path.exists() and not target.exists():
        copy2(path, target)

old = Presentation('/mnt/d/download/张异凡中期报告.pptx')
figure = max(
    (s for s in old.slides[9].shapes if s.shape_type == 13),
    key=lambda s: s.width * s.height,
)
(assets / 'depeweg_original.png').write_bytes(figure.image.blob)
with fitz.open('/home/erzhu419/mine_code/Catetorical-Feature-SAC-in-bus-simulation/Single Agent Robust Deep Reinforcement Learning for Bus Fleet Control.pdf') as paper:
    paper[6].get_pixmap(matrix=fitz.Matrix(5, 5), clip=fitz.Rect(155, 62, 470, 346)).save(assets / 'bus_figure2.png')
print('Backups preserved; original Figure 2 and midterm-slide illustration extracted.')
