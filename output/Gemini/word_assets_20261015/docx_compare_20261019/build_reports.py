"""Produce reviewable, complete difference tables without editing either DOCX."""
from pathlib import Path
from collections import Counter
from datetime import datetime
from zoneinfo import ZoneInfo
import json,re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

A=Path(__file__).resolve().parent
R=A.parents[1]
def read(n):return json.loads((A/n).read_text(encoding='utf-8-sig'))
n=read('native_comparison.json');c=read('content_independent.json')
m=read('math_figures_tables_independent.json');l=read('layout_package_independent.json')
assert all(d['baseline']==n['baseline'] for d in (c,m,l))
assert n['baseline'].startswith('/mnt/d/')
assert m['all_376_numbered_formula_content_and_structure_equal']
assert m['all_real_table_numeric_tokens_and_grid_equal']
assert c['summary']['bibliography_differences']==0

def flat(obj,path=''):
    if isinstance(obj,dict):
        out={}
        for k,v in obj.items():out.update(flat(v,path+'/'+str(k)))
        return out
    if isinstance(obj,list):
        out={}
        for i,v in enumerate(obj):out.update(flat(v,path+'/'+str(i+1)))
        return out
    return {path:obj}
def delta(a,b):
    aa,bb=flat(a),flat(b)
    return [(k,aa.get(k),bb.get(k)) for k in sorted(set(aa)|set(bb)) if aa.get(k)!=bb.get(k)]
def val(v):
    if v is None:return '未设置/不存在'
    if isinstance(v,(dict,list)):return json.dumps(v,ensure_ascii=False)
    return str(v)
def clean_changes(rows):
    return [r for r in rows if '/w:rsid/' not in r['property']]

wb=Workbook();wb.remove(wb.active);counts={}
def sheet(name,headers,rows,widths=None):
    ws=wb.create_sheet(name);ws.append(headers)
    for row_number,row in enumerate(rows,2):
        prepared=[val(v) for v in row]
        assert all(len(s)<=32767 for s in prepared),(name,[len(s) for s in prepared])
        ws.append(prepared)
        for column_number in range(1,len(headers)+1):
            cell=ws.cell(row_number,column_number)
            cell.data_type='s'
            cell.alignment=Alignment(vertical='top',wrap_text=True)
        ws.row_dimensions[row_number].height=48
    for cell in ws[1]:
        cell.font=Font(name='微软雅黑',bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='24476A')
        cell.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=30;ws.freeze_panes='A2'
    ws.auto_filter.ref=ws.dimensions
    for i in range(1,len(headers)+1):ws.column_dimensions[get_column_letter(i)].width=widths[i-1] if widths else (26 if i<=2 else 62)
    counts[name]=ws.max_row-1
    return ws

issues=[
 ('明确文字问题','3.5 本章小结','5.00 s；7.35 s','5.00" s"；7.35" s"','2处单位新增字面双引号，原生Word截图可见'),
 ('明确文字问题','5.5 本章小结','16.66 s；19.75 s','16.66" s"；19.75" s"','2处单位新增字面双引号'),
 ('明确文字问题','6.5 本章小结','51.14、69.93、87.05、57.67、70.67、76.26、4.03、44.19的单位s','对应12处问题中的8处变为" s"','8处单位新增字面双引号，数字未变'),
 ('明确文字问题','6.4.4 第2段','源仿真标定可达时以第5章跨保真混合训练为基础默认方案','源仿真标定可实现时以章跨保真混合训练为基础默认方案','漏掉“第5”；同段另删“第5章的”'),
 ('目录内容缺项','第4章目录','117条目录项，包括3个第4章四级条目','114条；该3条从当前目录消失','正文对应标题仍在'),
 ('引用格式改变','部分正文引文','上标','239个已匹配引用字符改成普通基线','239是字符数，不是239篇文献；链接目标仍有效'),
 ('行距改变','368个段落','20磅最小行距','20磅固定行距','含大量行内公式，原先可随高公式伸展的属性被移除'),
 ('表格分页改变','表5-6 更新次数严格匹配的三类训练方案配置','4行禁止跨页；首行重复表头','4行cantSplit和首行tblHeader删除','现允许行跨页，首行不再自动重复'),
 ('目录配置待确认','20个标题及TOC字段','自定义模板标题样式与TOC的\\t清单对应','标题改heading1/2/3，TOC清单未跟随修改','更新目录时需要核对是否仍完整收录；未据静态配置断言更新结果'),
 ('新增待确认内容','前置/后置页','没有这两页封面及成果、致谢页','新增2封面；6条成果另留空[7]；致谢“感谢导师，感谢国家。”','封面研究方向、日期、答辩等仍有空项'),
 ('其他真实文字改动','1.1、3.1、4.4','原段落措辞','删“固定”加“但是”；3.1调语序；4.4“厘清”→“理清”','完整旧新段落及逐字符变化见文字对照'),
 ('页眉布局改变','第3/4章','学校名上行左；完整章名下行右','学校名左；完整长章名右侧两行','实际显示完整，不能将XML串接顺序判为乱码'),
 ('图片版式改变','全部53图','原显示大小','53图外框均有变化，多数为微小取整；另有实质缩放','6图裁切均只移除白边，未切到图中文字/曲线'),
 ('公式对象改变','若干长式和16行内表达式','原生公式对象','12长式拆对象断行；16行内表达式转普通文字','376个编号式的数学内容、结构与编号顺序保持'),
 ('其他格式改变','表内段/图表题/字体字号/列表','原格式','83表内段改行距；图表题段前后改；部分字体字号及斜体改变；自动编号改文字编号','逐项旧新值见相应工作表'),
 ('保持正确','数据、引用、页面','376编号式；53图数据；47既有表；218文献；16节纸张和边距','均保持；引用目标齐全','新增2封面表不计入47既有表；未发现修订/批注/新增隐藏正文或嵌入程序部件'),
]
sheet('结论与重点差异',['类别','位置','基准版','调整版','说明'],issues,[21,35,65,65,75])

def content_kind(r):
    if r['kind']=='目录页码' and (not r['before'] or not r['after']):return '目录项增删'
    return r['kind']
sheet('文字与目录完整对照',['序号','类型','位置','基准正文块序号','调整正文块序号','基准全文','调整全文','逐字符变化'],[
 (r['id'],content_kind(r),r['location'],r['old_body_blocks'],r['new_body_blocks'],r['before'],r['after'],r['fragments'])for r in n['content']
],[9,22,50,20,20,95,95,85])
sheet('独立内容组对照',['序号','分类','位置','基准段范围','调整段范围','基准全文','调整全文','字符变化'],[
 (r['id'],r['kind'],r['section'],r['old_range_nonempty_1based'],r['new_range_nonempty_1based'],r['before'] if r['kind']=='content_change' else r['before_text'],r['after'] if r['kind']=='content_change' else r['after_text'],r.get('character_changes',[]))for r in c['diffs']
],[9,22,45,18,18,95,95,85])
sheet('376公式逐式核对',['公式号','内容结构相同','基准公式对象数','调整公式对象数','差异性质','基准可提取文字','调整可提取文字'],[
 (r['label'],r['content_and_math_structure_equal'],r['old_oMath_count'],r['new_oMath_count'],r.get('kind','保持'),r.get('old_text',''),r.get('new_text',''))for r in m['numbered_formula_xml_changes']
],[16,20,18,18,36,100,100])
conversions=[]
for r in m['all_math_stream_changes']:
    for x in r.get('conversions',[]):
        conversions.append((r.get('kind'),x['old_expression'],x['new_plain_expression'],x['extra_literal_unit_quotes'],x['old_paragraph'],'\n\n'.join(x['new_paragraphs'])))
sheet('16行内公式转文字',['类别','原公式文字','调整后普通文字','多出字面引号','所在原段全文','调整段全文'],conversions,[30,30,30,18,100,100])
crop={r['figure']:r for r in m['crop_content_bounds_check']['checks']}
figures=[]
for r in m['figure_differences']:
    ch=r['changes'];extent=ch.get('extent',{});f=r['figure']
    def size(which):
        a=extent.get(which,[])
        return '；'.join(f'{int(v["cx"])/360000:.4f}×{int(v["cy"])/360000:.4f}cm'for v in a)
    figures.append((f,r['image_bytes_equal'],r['caption_text_equal'],size('before'),size('after'),ch.get('crop',{}),crop.get(f,{}),ch))
sheet('53图片显示差异',['图号','图片字节不变','图题文字不变','基准显示宽高','调整显示宽高','裁切旧新值','裁切内容边界核对','全部显示属性差异'],figures,[15,18,18,26,26,55,65,100])
sheet('47既有表数据核对',['序号','表题','基准正文块0起算','调整正文块0起算','行列合并网格相同','文字逐字相同','去空白文字相同','数值逐单元格相同','变动单元格'],[
 (r['ordinal'],r['caption'],r['old_body_index'],r['new_body_index'],r['row_cell_grid_equal'],r['cell_text_exact_equal'],r['cell_text_equal_ignoring_whitespace'],r['numeric_tokens_per_cell_equal'],next((x['changed_cells']for x in m['table_differences']if x['ordinal']==r['ordinal']),[]))for r in m['table_checks']
],[9,65,20,20,20,20,20,20,80])

sheet('段落直接格式全部差异',['位置','基准地址','调整地址','基准文字','调整文字','全部旧新属性'],[
 (r['location'],r['old_address'],r['new_address'],r['text_before'],r['text_after'],clean_changes(r['changes'])) for r in n['paragraph_format'] if clean_changes(r['changes'])
],[50,28,28,75,75,100])
sheet('字符直接格式全部差异',['位置','基准地址','调整地址','基准字符区间0起算','调整字符区间0起算','对应文字','全部旧新属性'],[
 (r['location'],r['old_address'],r['new_address'],r['old_span'],r['new_span'],r['text'],clean_changes(r['changes']))for r in n['character_format']if clean_changes(r['changes'])
],[50,28,28,20,20,65,110])
sheet('段落有效格式变化',['基准段0起算','调整段0起算','原文定位','全部有效属性变化'],[
 (place['before_p'],place['after_p'],place['text'],group['differences'])for group in l['effective_format']['paragraph_change_groups']for place in group['locations']
],[20,20,85,110])
sheet('字符有效格式汇总',['属性','原值','新值','受影响已匹配字符数','示例定位'],[
 (r['property'],r['before'],r['after'],r['characters'],r['example_locations'])for r in l['effective_format']['run_character_transitions']
],[30,32,32,24,110])
sheet('样式定义全部差异',['样式名','基准ID','调整ID','基准使用段数','调整使用段数','全部属性变化'],[
 (r['style_name'],r['old_style_id'],r['new_style_id'],r['old_usage'],r['new_usage'],clean_changes(r['changes']))for r in n['style_definitions']if clean_changes(r['changes'])
],[45,25,25,20,20,110])
table_rows=[]
for r in n['table_format']:
    for prop,before,after in delta(r['before'],r['after']):table_rows.append((r['location'],r['old_address'],r['new_address'],prop,before,after))
sheet('表格格式属性逐项',['位置','基准地址','调整地址','属性路径','原值','新值'],table_rows,[55,25,25,95,40,40])
sheet('分节与页眉页脚',['类别','位置','原值','新值','差异说明'],[
 ('分节',str(r['section_index_1based']),{k:v['before']for k,v in r['differences'].items()},{k:v['after']for k,v in r['differences'].items()},r['differences'])for r in l['sections']['differences']
]+[
 ('页眉页脚',f'节{r["section_index_1based"]} {r["slot"]}',r['before_text'],r['after_text'],{'before_part':r['before_part'],'after_part':r['after_part']})for r in l['header_footer_changes']
],[25,50,100,100,100])
sheet('设置元数据包结构',['位置','项目','原值','新值'],[
 (section,key,v.get('before'),v.get('after'))for section in ('settings_changes','core_metadata','app_metadata','custom_metadata','web_settings')for key,v in l.get(section,{}).items()
]+[('包部件',r['part'],r['old_bytes'],str(r['new_bytes'])+' '+r['status'])for r in n['package_parts']
],[35,60,100,100])
sheet('引用与修订核对',['项目','基准','调整','结论'],[
 ('参考文献条数',218,218,'218条内容逐条不变'),
 ('REF/PAGEREF/内部超链接缺失目标',0,0,'目标有效'),
 ('章节编号标题缺失',0,0,'正文编号标题保持'),
 ('修订记录',c['hidden_and_revisions']['baseline']['revisions'],c['hidden_and_revisions']['candidate']['revisions'],'无修订记录'),
 ('批注部件',c['hidden_and_revisions']['baseline']['comment_parts'],c['hidden_and_revisions']['candidate']['comment_parts'],'无批注部件'),
 ('隐藏正文','无','无','未新增隐藏正文'),
 ('新增宏/OLE/ActiveX/altChunk','无','无','未发现上述新增嵌入对象'),
 ('MS Mincho在普通正文出现',0,0,'新增224处均在原生数学对象中'),
 ('原生Word当前分页',302,289,'本机不保存的排版观察；不是内容删除数量'),
 ('完整原始证据目录',str(A),str(A),'native_comparison.json与3份independent.json'),
],[45,70,70,90])
sheet('其余格式差异摘要',['分类','位置','变化','影响'],[(r['type'],r['location'],r['finding'],r['impact'])for r in l['concise_findings']],[40,60,100,100])

out=R/'Word差异逐项对照_20261019.xlsx';wb.save(out)
# Detect a report export that omits a whole category or truncates a sheet.
with_check=load_workbook(out,read_only=True,data_only=False)
assert all(with_check[k].max_row-1==v for k,v in counts.items())
assert counts['53图片显示差异']==53 and counts['376公式逐式核对']==376 and counts['47既有表数据核对']==47
with_check.close()

full=R/'Word文字差异全文_20261019.md'
lines=['# Word文字差异完整对照','',f'基准：`{n["baseline"]}`',f'待查：`{n["compared"]}`','',
 '方向固定为：基准版 → 调整版。以下完整列出所有提取文字变化及目录结果变化；自动编号、公式括号/分式结构、图片和格式另在Excel核对，不能仅凭提取文字判断数学等价。英文摘要原自动(1)–(4)改为文字编号，不属于新增四项研究结果。','']
for r in n['content']:
    lines += [f'## {r["id"]}. {content_kind(r)}｜{r["location"]}','',f'基准正文块：{r["old_body_blocks"]}；调整正文块：{r["new_body_blocks"]}。','',
     '**基准全文**','',r['before'] or '（无）','','**调整全文**','',r['after'] or '（无）','','**逐字符变化**','']
    for d in r['fragments']:lines += [f'- `{d["op"]}`：{json.dumps(d["before"],ensure_ascii=False)} → {json.dumps(d["after"],ensure_ascii=False)}']
    lines += ['']
full.write_text('\n'.join(lines)+'\n')

report=R/'Word差异核对说明_20261019.md'
lines=['# Word差异核对说明','',f'基准（用户最终明确D盘）：[{Path(n["baseline"]).name}]({n["baseline"]})。',f'待查：[{Path(n["compared"]).name}]({n["compared"]})。','',
 '两个原文件均未保存或改写。对方文件不只是格式调整，已发现明确文字问题和排版属性变化。没有据此判断对方的动机。','',
 f'[全部逐项差异 Excel]({out}) · [文字及目录完整旧新全文]({full})','',
 '| 类别 | 位置 | 具体变化 |','| --- | --- | --- |']
for typ,where,old,new,note in issues:
    lines.append('| '+typ+' | '+where+' | '+(old+' → '+new+'。'+note).replace('|','\\|')+' |')
lines += ['',
 '六张裁切图为3-6、3-7、3-12、4-4、4-7、5-2；对照原PNG内容边界，裁切只移除白边，没有切到图中文字或曲线。53张原图字节完全相同，53图题和46表题文字相同。47个既有表（含主要符号表）逐单元格数值与合并网格相同，表4-2只有两处空白变化；另新增两张封面信息表。','',
 '376个编号公式逐式核对通过，编号、顺序和数学结构相同；12个长式拆分为多个数学对象以换行。16个原生行内表达式转成普通文字，其中12个秒单位新增字面引号；这些引号已在Microsoft Word实际页面中看到。','',
 '目录漏项、上标和行距等的每处定位、全部直接格式旧新值、字体/字号有效变化、分节/页眉页脚与包结构变化均列于Excel。内建样式ID重命名、WPS编辑元数据和字段缓存更新单独归类，不当成文字丢失。','',
 '## 核对范围与限制','',
 '已逐项比较完整OOXML文字、376编号公式、53图像、47既有表、218参考文献、引用目标、样式与直接格式。原生Word只查看了选定页面，未逐页肉眼检查全部302/289页；行距保护被移除不等于已证明每个公式都被截断。目录样式配置的后续更新结果需在副本中确认。本次不修改论文、不导出PDF。','',
 '## 原始证据','',
 f'- [完整文本与格式属性差异]({A/"native_comparison.json"})',
 f'- [独立文字和引用核对]({A/"content_independent.json"})',
 f'- [公式、图片、表格核对]({A/"math_figures_tables_independent.json"})',
 f'- [有效格式与包结构核对]({A/"layout_package_independent.json"})',
 f'- [Word实际页：单位引号]({A/"adjusted_page_112.png"})',
 f'- [同段基准页]({A/"baseline_page_117.png"})',
]
report.write_text('\n'.join(lines)+'\n')
(A/'report_export.json').write_text(json.dumps({'baseline':n['baseline'],'compared':n['compared'],'excel':str(out),'report':str(report),'text_full':str(full),'sheet_rows':counts,'original_documents_modified':False},ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'excel':str(out),'report':str(report),'text_full':str(full),'sheet_rows':counts},ensure_ascii=False))
