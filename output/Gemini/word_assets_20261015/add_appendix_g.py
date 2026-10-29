"""Retain original appendix F and append the latest supplementary settings as G."""
from copy import deepcopy
from pathlib import Path
import json,re,shutil
import pypandoc
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path('/home/erzhu419/mine_code/Thesis')
A=Path(__file__).resolve().parent
TARGET=ROOT/'output/Gemini/重构论文_排版稿_20261015.docx'
BACKUP=A/'附录G调整前.docx'
if not BACKUP.exists():shutil.copy2(TARGET,BACKUP)
doc=Document(BACKUP)
def text(e):return ''.join(n.text or '' for n in e.iter() if n.tag in [qn('w:t'),qn('m:t')])
def sty(e):
 n=e.find('./'+qn('w:pPr')+'/'+qn('w:pStyle'))
 return n.get(qn('w:val')) if n is not None else ''
def han(s):return ''.join(re.findall(r'[\u3400-\u9fff]',s))
def para(s='',style='body'):
 e=OxmlElement('w:p');p=Paragraph(e,doc._body);p.style=doc.styles.get_by_id(style,1)
 if s:p.add_run(s)
 return e

source=(ROOT/'output/Gemini/重构论文.md').read_text().split('\n# 附录F',1)[1]
source='# 附录G'+source
source=re.sub(r'^---[ \t]*$','',source,flags=re.M).strip()
source=source.replace('\\tag{附F-','\\tag{附G-')
original_g_source=source
# All Chinese is retained; three body references change only the appendix letter.
changed=0
for p in doc.element.iter(qn('w:p')):
 ts=list(p.iter(qn('w:t')))
 value=''.join(n.text or '' for n in ts)
 for match in re.finditer('附录F第五部分',value):
  pos=match.start()+2;offset=0
  for n in ts:
   v=n.text or ''
   if offset<=pos<offset+len(v):
    local=pos-offset;n.text=v[:local]+'G'+v[local+1:];changed+=1;break
   offset+=len(v)
assert changed==3,changed

combined=json.loads((A/'citation_key_map.json').read_text())['matched']
for p in A.glob('citation_additions_*.json'):
 if p.name.endswith('_unresolved.json'):continue
 d=json.loads(p.read_text());combined.update(d.get('matched',d))
numbers=json.loads((A/'reference_numbers.json').read_text())
body=list(doc.element.body)
bib=[e for e in body if sty(e)=='bibliography']
last_bib=bib[-1]
maxnum=max(int(re.match(r'\[(\d+)\]',text(e))[1]) for e in bib)
nextmark=max(int(n.get(qn('w:id'))) for n in doc.element.iter(qn('w:bookmarkStart')))+100
added=[]

def citation(m):
 global maxnum,nextmark,last_bib
 refs=[]
 for key in [k.strip() for k in m[1].split(',')]:
  if key not in numbers:
   rec=combined[key];oldid=rec.get('old_id',key)
   alias=next((k for k in numbers if k in combined and combined[k].get('old_id',k)==oldid),None)
   if alias:numbers[key]=numbers[alias]
   else:
    maxnum+=1;numbers[key]=maxnum
    entry=re.sub(r'^\[\d+\]\s*','',rec['bibliography_text'])
    e=para('['+str(maxnum)+']\t'+entry,'bibliography');p=Paragraph(e,doc._body)
    p.paragraph_format.left_indent=Pt(30);p.paragraph_format.first_line_indent=Pt(-30);p.paragraph_format.tab_stops.add_tab_stop(Pt(30))
    b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(nextmark));b.set(qn('w:name'),'latest_ref_'+str(maxnum));e.insert(1,b)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(nextmark));e.append(b);nextmark+=1
    last_bib.addnext(e);last_bib=e;added.append(key)
  refs.append('['+str(numbers[key])+'](#latest_ref_'+str(numbers[key])+')')
 return '^\\['+','.join(refs)+'\\]^'
source=re.sub(r'\[([A-Za-z]+\d{4}[\w-]*(?:\s*,\s*[A-Za-z]+\d{4}[\w-]*)*)\]',citation,source)
tags=[]
def equation(m):
 formula=m[1];tag=re.search(r'\\tag\{([^}]+)\}',formula)[1];tags.append(tag)
 formula=re.sub(r'\\tag\{[^}]+\}','',formula).strip()
 # The first formula contains three independent identities, set on three lines.
 if tag=='附G-1':
  rows=formula.split('\\qquad')
  formula='\\begin{aligned}'+r'\\'.join(row.strip().replace(r'\;=\;',r'&=',1) for row in rows)+'\\end{aligned}'
 return '\n\n$$'+formula+'$$\n\nAPPENDIXGNUMBER'+str(len(tags))+'END\n\n'
# Do not alter formula semantics; ampersands above are alignment markers.
source=re.sub(r'\$\$(.*?)\$\$',equation,source,flags=re.S)
source=source.replace(r'\u0026','&')
converted=A/'appendix_g_native.docx'
pypandoc.convert_text(source,'docx',format='markdown+tex_math_dollars-smart',outputfile=str(converted),extra_args=['--reference-doc='+str(BACKUP)])
new=Document(converted)
newstyles={s.style_id:s.name for s in new.styles}
nodes=[]
width=Emu(doc.sections[-1].page_width-doc.sections[-1].left_margin-doc.sections[-1].right_margin)
for raw in new.element.body:
 if raw.tag==qn('w:sectPr'):continue
 e=deepcopy(raw)
 if e.tag!=qn('w:p'):
  for mark in e.iter():
   if mark.tag in [qn('w:bookmarkStart'),qn('w:bookmarkEnd')]:mark.set(qn('w:id'),str(int(mark.get(qn('w:id')))+nextmark))
  nodes.append(e);continue
 p=Paragraph(e,doc._body);value=text(e)
 if value.startswith('APPENDIXGNUMBER'):
  idx=int(re.fullmatch(r'APPENDIXGNUMBER(\d+)END',value)[1])-1
  formula=nodes.pop();assert formula.find('.//'+qn('m:oMath')) is not None
  table=doc.add_table(rows=1,cols=3);table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
  for cell,col,cw in zip(table.rows[0].cells,table.columns,[Pt(43),width-Pt(86),Pt(43)]):
   cell.width=col.width=cw;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    n=OxmlElement('w:'+side);n.set(qn('w:w'),'0');n.set(qn('w:type'),'dxa');mar.append(n)
   cell._tc.get_or_add_tcPr().append(mar)
   for pp in cell.paragraphs:
    pp.style=doc.styles.get_by_id('body',1);pp._p.get_or_add_pPr().get_or_add_ind().set(qn('w:firstLineChars'),'0');pp.paragraph_format.first_line_indent=Pt(0);pp.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST;pp.paragraph_format.line_spacing=Pt(20)
  center=table.cell(0,1);center._tc.remove(center.paragraphs[0]._p);center._tc.append(formula)
  fp=Paragraph(formula,doc._body);fp.style=doc.styles.get_by_id('body',1);fp._p.get_or_add_pPr().get_or_add_ind().set(qn('w:firstLineChars'),'0');fp.paragraph_format.first_line_indent=Pt(0);fp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
  rp=table.cell(0,2).paragraphs[0];rp.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.RIGHT;rp.add_run('('+tags[idx]+')').font.size=Pt(10.5)
  borders=OxmlElement('w:tblBorders')
  for side in ['top','bottom','left','right','insideH','insideV']:
   n=OxmlElement('w:'+side);n.set(qn('w:val'),'nil');borders.append(n)
  table._tbl.tblPr.append(borders);table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  table._tbl.getparent().remove(table._tbl);nodes.append(table._tbl);continue
 if newstyles.get(sty(e),'')=='Heading 1':p.style=doc.styles.get_by_id('chapter',1)
 else:p.style=doc.styles.get_by_id('body',1)
 if e.find('.//'+qn('m:oMath')) is not None:p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST;p.paragraph_format.line_spacing=Pt(20)
 for n in e.iter():
  if n.tag in [qn('w:bookmarkStart'),qn('w:bookmarkEnd')]:n.set(qn('w:id'),str(int(n.get(qn('w:id')))+nextmark))
  if n.tag==qn('w:bookmarkStart'):n.set(qn('w:name'),'appendix_g_'+n.get(qn('w:name')))
  if n.tag==qn('w:hyperlink') and n.get(qn('w:anchor')):
   assert n.get(qn('w:anchor')).startswith('latest_ref_')
   for run in n.iter(qn('w:r')):
    pr=run.find(qn('w:rPr'))
    if pr is None:pr=OxmlElement('w:rPr');run.insert(0,pr)
    c=OxmlElement('w:color');c.set(qn('w:val'),'000000');pr.append(c)
    u=OxmlElement('w:u');u.set(qn('w:val'),'none');pr.append(u)
 nodes.append(e)
assert han(''.join(text(e) for e in nodes))==han(original_g_source),'G Chinese changed'
# End F with its original section settings, then create a new section for G.
sect=doc.element.body.sectPr
oldsect=deepcopy(sect);doc.element.body.remove(sect)
breakp=para();breakp.find(qn('w:pPr')).append(oldsect);doc.element.body.append(breakp)
for e in nodes:doc.element.body.append(e)
gsect=deepcopy(oldsect)
for n in list(gsect):
 if n.tag in [qn('w:headerReference'),qn('w:footerReference')]:gsect.remove(n)
doc.element.body.append(gsect)
original_header=deepcopy(doc.sections[-2].header.paragraphs[0]._p)
s=doc.sections[-1];s.header.is_linked_to_previous=False
h=s.header.paragraphs[0]
pr=h._p.find(qn('w:pPr'))
if pr is not None:h._p.remove(pr)
pr=original_header.find(qn('w:pPr'))
if pr is not None:h._p.insert(0,deepcopy(pr))
# Clone the original header typography and tabs, changing its chapter text only.
h.text='中南大学博士学位论文\t附录G 受控动力学任务模型与跨保真补充实验设置'
for r in h.runs:r.font.size=Pt(9)
doc.save(TARGET)
(A/'reference_numbers.json').write_text(json.dumps(numbers,ensure_ascii=False,indent=2))
result=json.loads((A/'build_result.json').read_text());result.update(annexes='Original Word A-F plus latest manuscript G',new_appendix_g_chinese_unchanged=True,new_appendix_g_display_equations=len(tags),body_appendix_reference_letters_updated=changed,bibliography_count=maxnum,native_math=len(doc.element.xpath('.//m:oMath')),word_finalized=False)
(A/'build_result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps({'output':str(TARGET),'G_paragraphs':len(nodes),'G_equations':tags,'new_bibliography':added,'bibliography_total':maxnum,'body_references_changed':changed},ensure_ascii=False,indent=2))
