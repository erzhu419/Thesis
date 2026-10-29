"""Add editable before/after bus disruption diagrams to page 19 only."""
import json
from pathlib import Path
import posixpath
from shutil import copy2
import sys
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / 'workpapers/博士答辩_20261007'))
from drawing import Deck, TEAL, RED, GOLD, INK, MUTED, LINE, LIGHT, WHITE

SOURCE = ROOT / 'output/博士毕业答辩_公交原理补充_20261011.pptx'
OUTPUT = ROOT / 'output/博士毕业答辩_突变前后补图_20261011.pptx'


def move(shape, y, h, size=None):
    shape.top, shape.height = Inches(y), Inches(h)
    if size is not None:
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(size)


def person(deck, x, y, color):
    deck.ellipse(x+.05, y, .085, .085, color)
    deck.line(x+.092, y+.1, x+.092, y+.24, color, 1.8)
    deck.line(x+.092, y+.14, x+.02, y+.2, color, 1.8)
    deck.line(x+.092, y+.14, x+.17, y+.2, color, 1.8)
    deck.line(x+.092, y+.24, x+.03, y+.34, color, 1.8)
    deck.line(x+.092, y+.24, x+.16, y+.34, color, 1.8)


def stop(deck, x):
    deck.rect(x, 3.36, .44, .2, LIGHT, TEAL)
    deck.text(x, 3.36, .44, .2, '站点', 10, True, TEAL, PP_ALIGN.CENTER)
    deck.line(x+.22, 3.56, x+.22, 3.76, TEAL, 1.5)


def demand(deck, x, surge):
    color = RED if surge else TEAL
    deck.text(x, 3.0, 5.38, .34, '突发客流激增' if surge else '客流平稳', 17, True, color)
    deck.line(x+.02, 3.76, x+5.25, 3.76, MUTED, 1.6)
    stop(deck, x+.07)
    xs = [x+.85+j*.33 for j in range(9)] if surge else [x+.88, x+1.42, x+1.96]
    for xx in xs:
        person(deck, xx, 3.38, color)
    deck.bus(x+4.07, 3.4, .65, color=TEAL)
    deck.text(x+.07, 3.85, 5.14, .27,
              '站台队列拉长，上下车服务时间增加' if surge else '短队列，按原有节奏服务',
              15, color=color)


def junction(deck, x, closure):
    color = RED if closure else TEAL
    deck.text(x, 4.15, 5.38, .34, '路口临时交通管制' if closure else '路口正常通行', 17, True, color)
    deck.rect(x+.03, 4.58, 5.2, .3, 'E4EAEC', None)
    deck.rect(x+3.77, 4.46, .39, .49, 'E4EAEC', None)
    deck.line(x+.03, 4.72, x+5.23, 4.72, WHITE, .7, dashed=True)
    signal_x = x+3.39
    deck.rect(signal_x, 4.44, .16, .3, INK, None, True)
    deck.ellipse(signal_x+.04, 4.485 if closure else 4.63, .08, .08, RED if closure else TEAL)
    deck.line(signal_x+.08, 4.74, signal_x+.08, 4.89, INK, 1.3)
    deck.bus(x+1.14, 4.53, .59, color=TEAL)
    if closure:
        deck.bus(x+.29, 4.53, .59, color=TEAL)
        barrier_x = x+3.66
        deck.rect(barrier_x, 4.61, 1.0, .12, RED, None)
        for j in range(4):
            deck.rect(barrier_x+.09+j*.22, 4.61, .1, .12, WHITE, None)
        deck.line(barrier_x+.07, 4.73, barrier_x+.07, 4.89, RED, 1.5)
        deck.line(barrier_x+.93, 4.73, barrier_x+.93, 4.89, RED, 1.5)
    else:
        deck.line(x+2.0, 4.7, x+2.9, 4.7, TEAL, 2, arrow=True)
    deck.text(x+.07, 4.99, 5.14, .26,
              '通行中断，延误骤增，车辆积压' if closure else '正常通过，运行时间保持原有规律',
              14, color=color)


def main():
    backup = HERE / 'before_shock_comparison'
    backup.mkdir(exist_ok=True)
    if not (backup / SOURCE.name).exists():
        copy2(SOURCE, backup / SOURCE.name)
    deck = Deck()
    deck.prs = Presentation(SOURCE)
    slide = deck.prs.slides[18]
    deck.slide = slide
    original_shapes = list(slide.shapes)
    by_id = {shape.shape_id:shape for shape in original_shapes}
    move(by_id[9], 1.79, .38, 19)
    move(by_id[10], 2.14, .26, 12)
    for bg in [11,15,19]:
        move(by_id[bg], 5.46, .62)
    for title in [12,16,20]:
        move(by_id[title], 5.49, .28, 16)
    for sub in [13,17,21]:
        move(by_id[sub], 5.8, .25, 13)
    for arrow in [14,18]:
        by_id[arrow].top = Inches(5.75)

    deck.text(.91, 2.55, 5.41, .41, '突变前', 21, True, TEAL, PP_ALIGN.CENTER)
    deck.text(7.01, 2.55, 5.41, .41, '突变后', 21, True, RED, PP_ALIGN.CENTER)
    deck.line(6.68, 2.55, 6.68, 5.18, LINE, 1)
    demand(deck, .93, False)
    demand(deck, 7.03, True)
    junction(deck, .93, False)
    junction(deck, 7.03, True)
    for shape in list(slide.shapes)[len(original_shapes):]:
        for effect in shape._element.xpath('.//a:effectRef'):
            effect.set('idx', '0')
    notes = slide.notes_slide.notes_text_frame
    lines = notes.text.splitlines()
    lines.insert(1, '新增图为突变前后场景示意：客流激增使站台队列和上下车服务时间增加；临时交通管制使通行中断、延误骤增并造成车辆积压。两类变化都要求发现变化后合理调整驻站响应。图中人数与车辆仅解释场景，不是新增实验。')
    notes.text = '\n'.join(lines)
    deck.prs.save(OUTPUT)

    # Preserve every original package part except this slide and its notes.
    with ZipFile(SOURCE) as original:
        presentation = etree.fromstring(original.read('ppt/presentation.xml'))
        rels = etree.fromstring(original.read('ppt/_rels/presentation.xml.rels'))
        targets = {rel.get('Id'):rel.get('Target') for rel in rels}
        ids = presentation.find('{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst')
        rid = ids[18].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        slide_part = posixpath.normpath('ppt/'+targets[rid])
    with ZipFile(OUTPUT) as edited:
        replacements = {
            slide_part:edited.read(str(slide.part.partname).lstrip('/')),
            str(slide.notes_slide.part.partname).lstrip('/'):edited.read(str(slide.notes_slide.part.partname).lstrip('/')),
        }
    with ZipFile(SOURCE) as original, ZipFile(OUTPUT,'w') as output:
        for item in original.infolist():
            output.writestr(item, replacements.get(item.filename,original.read(item.filename)))
    check = Presentation(OUTPUT)
    before = Presentation(SOURCE)
    assert len(check.slides) == len(before.slides) == 57
    assert all(etree.tostring(a._element)==etree.tostring(b._element) for i,(a,b) in enumerate(zip(before.slides,check.slides)) if i!=18)
    texts = {s.shape_id:s.text for s in check.slides[18].shapes if s.has_text_frame}
    assert all(texts[s.shape_id]==s.text for s in before.slides[18].shapes if s.has_text_frame)
    assert SOURCE.read_bytes() == (backup/SOURCE.name).read_bytes()
    print(json.dumps({'slides':57,'modified_page':19,'existing_text_preserved':True,
                      'other_slides_preserved':True,'all_existing_charts_preserved':True},ensure_ascii=False))


if __name__ == '__main__':
    main()
