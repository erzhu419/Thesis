"""Number editable equations and connect thesis references in Word."""

import re
from copy import deepcopy

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

EQUATION = re.compile(r"(?ms)(?:<!-- eq:(\w+) -->\n\n)?(^\$\$[ \t]*\n.*?^\$\$[ \t]*$)")


def label_sources(paths):
    for path in paths:
        chapter = path.name[0] if path.parent.name == "appendices" else str(int(path.name[:2]))
        count = 0
        def replace(match):
            nonlocal count
            count += 1
            label = match[1] or f"eq{chapter}_{count:03}"
            return f"<!-- eq:{label} -->\n\n{match[2]}"
        path.write_text(EQUATION.sub(replace, path.read_text()))


def prepare_text(content):
    records, numbers, chapter_counts = [], {}, {}
    chapter = None
    for line in content.splitlines():
        heading = re.match(r"^# (?:第([1-7])章|附录([A-F]))", line)
        if heading:
            chapter = heading[1] or heading[2]
        match = re.fullmatch(r"<!-- eq:(\w+) -->", line)
        if match:
            assert chapter is not None, line
            chapter_counts[chapter] = chapter_counts.get(chapter, 0) + 1
            label = match[1]
            assert label not in numbers, label
            number = f"{chapter}-{chapter_counts[chapter]}"
            numbers[label] = number
            records.append(dict(label=label, chapter=chapter, ordinal=chapter_counts[chapter], number=number))
    assert len(records) == len(EQUATION.findall(content)), "Unlabelled display equation"
    text = re.sub(r"<!-- eq:\w+ -->\n\n", "", content)
    def eqref(match):
        label = match[1]
        assert label in numbers, f"Undefined equation: {label}"
        return f"[式（{numbers[label]}）](#{label})"
    text = re.sub(r"\[式\]\(#(eq\w+)\)", eqref, text)
    lines = []
    targets = {}
    for line in text.splitlines():
        caption = re.match(r"^([图表][1-7A-F]-\d+)\s+", line)
        if caption:
            targets[caption[1]] = ("fig" if caption[1][0] == "图" else "tbl") + caption[1][1:].replace("-", "_")
    for line in text.splitlines():
        if not re.match(r"^#|^[图表][1-7A-F]-\d+\s+", line):
            def figref(match):
                label = match[0]
                assert label in targets, f"Undefined figure or table: {label}"
                return f"[{label}](#{targets[label]})"
            line = re.sub(r"[图表][1-7A-F]-\d+", figref, line)
            line = re.sub(r"附录([A-F])(?![A-Za-z])", lambda m: f"[附录{m[1]}](#app{m[1]})", line)
            if not re.match(r"^\*\*(?:定理|引理|命题|推论)[A-E]\.\d+", line):
                line = re.sub(r"(定理|引理|命题|推论)([A-E])\.(\d+)",
                              lambda m: f"[{m[0]}](#th_{m[2]}_{m[3]}_{dict(定理='t',引理='l',命题='p',推论='c')[m[1]]})", line)
        lines.append(line)
    return "\n".join(lines), records, targets


def field(paragraph, instruction, result):
    node = OxmlElement("w:fldSimple")
    node.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = result
    run.append(text)
    node.append(run)
    paragraph._p.append(node)
    return node


def bookmark_prefix(paragraph, prefix, name, number):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    for run in paragraph.runs:
        if run.text.startswith(prefix):
            tail = deepcopy(run._r)
            run.text, remainder = prefix, run.text[len(prefix):]
            run._r.addprevious(start)
            run._r.addnext(end)
            if remainder:
                for child in list(tail):
                    if child.tag != qn("w:rPr"):
                        tail.remove(child)
                node = OxmlElement("w:t")
                node.set(qn("xml:space"), "preserve")
                node.text = remainder
                tail.append(node)
                end.addnext(tail)
            return
    raise ValueError(f"Cannot locate bookmark prefix {prefix!r}")


def add_numbering(doc, records, caption_targets):
    bookmark_id = max([int(x.get(qn("w:id"))) for x in doc.element.xpath(".//w:bookmarkStart")] + [0]) + 1
    for paragraph in list(doc.paragraphs):
        for label, name in caption_targets.items():
            if paragraph.text.startswith(label + " "):
                bookmark_prefix(paragraph, label, name, bookmark_id)
                bookmark_id += 1
                break
        appendix = re.match(r"^(附录[A-F])\s", paragraph.text)
        theorem = re.match(r"^((定理|引理|命题|推论)([A-E])\.(\d+))。", paragraph.text)
        if appendix:
            bookmark_prefix(paragraph, appendix[1], "app"+appendix[1][-1], bookmark_id)
            bookmark_id += 1
        elif theorem:
            suffix = dict(定理="t", 引理="l", 命题="p", 推论="c")[theorem[2]]
            bookmark_prefix(paragraph, theorem[1], f"th_{theorem[3]}_{theorem[4]}_{suffix}", bookmark_id)
            bookmark_id += 1
    equations = [p for p in doc.paragraphs if p._p.xpath("./m:oMathPara")]
    assert len(equations) == len(records), (len(equations), len(records))
    width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for paragraph, record in zip(equations, records):
        table = doc.add_table(rows=1, cols=3)
        table.autofit = False
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        # Equal side cells center the native display equation independently of its number.
        widths = (Pt(34), width - Pt(68), Pt(34))
        for column, cell, cell_width in zip(table.columns, table.rows[0].cells, widths):
            column.width = cell.width = cell_width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                child = OxmlElement("w:"+side)
                child.set(qn("w:w"), "0")
                child.set(qn("w:type"), "dxa")
                margins.append(child)
            cell._tc.get_or_add_tcPr().append(margins)
            pf = cell.paragraphs[0].paragraph_format
            pf.first_line_indent = Pt(0)
            pf.line_spacing = Pt(20)
            pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            child = OxmlElement("w:"+side)
            child.set(qn("w:val"), "nil")
            borders.append(child)
        table._tbl.tblPr.append(borders)
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        paragraph._p.addprevious(table._tbl)
        middle = table.cell(0, 1)
        middle._tc.remove(middle.paragraphs[0]._p)
        middle._tc.append(paragraph._p)
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        right = table.cell(0, 2).paragraphs[0]
        right.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(bookmark_id)); start.set(qn("w:name"), record["label"])
        end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(bookmark_id))
        bookmark_id += 1
        right._p.append(start)
        right.add_run(f"({record['chapter']}-")
        field(right, f"SEQ Eq{record['chapter']} \\* ARABIC", str(record["ordinal"]))
        right.add_run(")")
        right._p.append(end)
        for run in right.runs:
            run.font.size = Pt(10.5)
    # REF fields remain updateable when the author edits equation numbering in Word.
    names = {x.get(qn("w:name")) for x in doc.element.xpath(".//w:bookmarkStart")}
    for link in list(doc.element.xpath(".//w:hyperlink[@w:anchor]")):
        name = link.get(qn("w:anchor"))
        if not name.startswith(("eq", "fig", "tbl", "app", "th_")):
            continue
        assert name in names, f"Missing Word reference target: {name}"
        cached = "".join(link.xpath(".//w:t/text()"))
        node = OxmlElement("w:fldSimple"); node.set(qn("w:instr"), f"REF {name} \\h")
        if name.startswith("eq"):
            cached = cached.removeprefix("式").replace("（", "(").replace("）", ")")
            prefix = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = "式"; prefix.append(text)
            link.addprevious(prefix)
        run = OxmlElement("w:r"); text = OxmlElement("w:t"); text.text = cached; run.append(text); node.append(run)
        link.addprevious(node); link.getparent().remove(link)
    return len(equations)
